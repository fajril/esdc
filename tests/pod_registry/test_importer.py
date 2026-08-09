"""Tests for the POD importer."""

# tests/pod_registry/test_importer.py
from datetime import datetime

import openpyxl
import pytest

import esdc.configs as configs
from esdc.pod_registry.importer import (
    PodRegistryImportError,
    import_pod_registry_workbook,
)
from esdc.pod_registry.store import get_sqlite_connection


def _patch_dirs(monkeypatch, tmp_path):
    monkeypatch.setattr(configs.Config, "get_db_dir", classmethod(lambda cls: tmp_path))
    monkeypatch.setattr(
        configs.Config, "get_db_file", classmethod(lambda cls: tmp_path / "esdc.duckdb")
    )


def _write_workbook(
    path, pod_rows=None, project_rows=None, revision_rows=None, institution_rows=None
):
    wb = openpyxl.Workbook()
    active = wb.active
    assert active is not None
    wb.remove(active)

    ws = wb.create_sheet("POD Record")
    ws.append(
        [
            "pod_id_itb",
            "approval_date",
            "institution",
            "pod_type",
            "rev_ num",
            "pod_name",
            "pod_letter_num",
            "approval_seq",
            "pod_id_skk",
            "preceded_by",
            "superseded_by",
        ]
    )
    for row in (
        pod_rows
        if pod_rows is not None
        else [
            [
                645,
                datetime(2003, 11, 21),
                "BP Migas",
                "POD/Waterflood/EOR",
                0,
                "POD Mengoepeh",
                "294/BP",
                5,
                "PL-2003-0005-3-2-0",
                None,
                "PL-2005-0051-3-2-1",
            ],
            [
                700,
                datetime(2005, 6, 1),
                "BP Migas",
                "POD/Waterflood/EOR",
                1,
                "POD Mengoepeh Rev",
                "51/BP",
                51,
                "PL-2005-0051-3-2-1",
                "PL-2003-0005-3-2-0",
                None,
            ],
        ]
    ):
        ws.append(row)

    ws = wb.create_sheet("project_pod")
    ws.append(["pod_id", "project_id"])
    for row in project_rows if project_rows is not None else [[645, "P-2403431-01"]]:
        ws.append(row)

    ws = wb.create_sheet("pod_revision")
    ws.append(["rev_id", "successor_id", "predecessor_id"])
    for row in (
        revision_rows
        if revision_rows is not None
        else [[1, "PL-2005-0051-3-2-1", "PL-2003-0005-3-2-0"]]
    ):
        ws.append(row)

    ws = wb.create_sheet("institution")
    ws.append(["code", "institution", "description"])
    for row in (
        institution_rows
        if institution_rows is not None
        else [[3, "BP Migas", "POD 2003 to 2012"], [4, "SKK Migas", "POD 2013 onward"]]
    ):
        ws.append(row)

    ws = wb.create_sheet("pod_type")
    ws.append(["code", "pod type", "description"])
    for row in [[1, "POD I", "the first POD"], [2, "POD/Waterflood/EOR", "other"]]:
        ws.append(row)

    wb.save(path)
    return path


def test_import_seeds_all_tables(monkeypatch, tmp_path):
    _patch_dirs(monkeypatch, tmp_path)
    xlsx = _write_workbook(tmp_path / "pod.xlsx")
    counts = import_pod_registry_workbook(xlsx)
    assert counts == {
        "r_institution": 2,
        "r_pod_type": 2,
        "m_pod": 2,
        "project_pod": 1,
        "pod_revision": 1,
    }
    conn = get_sqlite_connection()
    try:
        row = conn.execute("SELECT * FROM m_pod WHERE id = 645").fetchone()
        assert row["pod_id"] == "PL-2003-0005-3-2-0"
        assert row["approval_date"] == "2003-11-21"
        assert row["institution_code"] == 3
        assert row["pod_type_code"] == 2
        assert row["rev_num"] == 0
        assert row["approval_seq"] == 5
    finally:
        conn.close()


def test_import_is_idempotent(monkeypatch, tmp_path):
    _patch_dirs(monkeypatch, tmp_path)
    xlsx = _write_workbook(tmp_path / "pod.xlsx")
    import_pod_registry_workbook(xlsx)
    import_pod_registry_workbook(xlsx)
    conn = get_sqlite_connection()
    try:
        assert conn.execute("SELECT COUNT(*) FROM m_pod").fetchone()[0] == 2
    finally:
        conn.close()


def test_import_rejects_unknown_institution(monkeypatch, tmp_path):
    _patch_dirs(monkeypatch, tmp_path)
    xlsx = _write_workbook(
        tmp_path / "pod.xlsx",
        pod_rows=[
            [
                1,
                datetime(2003, 1, 1),
                "Unknown Body",
                "POD I",
                0,
                "X",
                None,
                1,
                "PL-2003-0001-1-1-0",
                None,
                None,
            ]
        ],
        project_rows=[],
        revision_rows=[],
    )
    with pytest.raises(PodRegistryImportError) as exc:
        import_pod_registry_workbook(xlsx)
    assert any("Unknown Body" in e for e in exc.value.errors)


def test_import_rejects_orphan_project_link(monkeypatch, tmp_path):
    _patch_dirs(monkeypatch, tmp_path)
    xlsx = _write_workbook(tmp_path / "pod.xlsx", project_rows=[[999, "P-X"]])
    with pytest.raises(PodRegistryImportError):
        import_pod_registry_workbook(xlsx)


def test_import_rejects_orphan_revision(monkeypatch, tmp_path):
    _patch_dirs(monkeypatch, tmp_path)
    xlsx = _write_workbook(
        tmp_path / "pod.xlsx",
        revision_rows=[[1, "PL-9999-9999-9-9-9", "PL-2003-0005-3-2-0"]],
    )
    with pytest.raises(PodRegistryImportError):
        import_pod_registry_workbook(xlsx)


def test_import_rejects_self_referencing_revision(monkeypatch, tmp_path):
    _patch_dirs(monkeypatch, tmp_path)
    xlsx = _write_workbook(
        tmp_path / "pod.xlsx",
        revision_rows=[[1, "PL-2003-0005-3-2-0", "PL-2003-0005-3-2-0"]],
    )
    with pytest.raises(PodRegistryImportError) as exc:
        import_pod_registry_workbook(xlsx)
    assert any("self-referencing" in e for e in exc.value.errors)


def test_import_rejects_missing_approval_seq(monkeypatch, tmp_path):
    _patch_dirs(monkeypatch, tmp_path)
    xlsx = _write_workbook(
        tmp_path / "pod.xlsx",
        pod_rows=[
            [
                645,
                datetime(2003, 11, 21),
                "BP Migas",
                "POD/Waterflood/EOR",
                0,
                "POD Mengoepeh",
                "294/BP",
                None,
                "PL-2003-0005-3-2-0",
                None,
                None,
            ]
        ],
        project_rows=[],
        revision_rows=[],
    )
    with pytest.raises(PodRegistryImportError) as exc:
        import_pod_registry_workbook(xlsx)
    assert any("approval_seq" in e for e in exc.value.errors)


def test_import_rejects_blank_project_pod_id(monkeypatch, tmp_path):
    """A blank pod_id cell in project_pod must surface as a row error, not crash.

    Bare int(r["pod_id"]) on a blank cell raises TypeError outside any
    try/except, which the CLI doesn't catch (only PodRegistryImportError) --
    the row error naming the sheet+row is the fix, and the good sibling row
    (645) must not also be flagged.
    """
    _patch_dirs(monkeypatch, tmp_path)
    xlsx = _write_workbook(
        tmp_path / "pod.xlsx",
        project_rows=[[None, "P-BLANK"], [645, "P-2403431-01"]],
    )
    with pytest.raises(PodRegistryImportError) as exc:
        import_pod_registry_workbook(xlsx)
    assert exc.value.errors == ["project_pod row 2: invalid pod_id 'None'"]


def test_import_rejects_text_project_pod_id(monkeypatch, tmp_path):
    """A non-numeric text pod_id cell must surface as a row error, not crash.

    Bare int(r["pod_id"]) on a text cell raises ValueError outside any
    try/except -- same fix as the blank-cell case.
    """
    _patch_dirs(monkeypatch, tmp_path)
    xlsx = _write_workbook(
        tmp_path / "pod.xlsx",
        project_rows=[["abc", "P-TEXT"], [645, "P-2403431-01"]],
    )
    with pytest.raises(PodRegistryImportError) as exc:
        import_pod_registry_workbook(xlsx)
    assert exc.value.errors == ["project_pod row 2: invalid pod_id 'abc'"]


def test_import_rejects_non_numeric_institution_code(monkeypatch, tmp_path):
    """A non-numeric institution code cell must surface as a row error, not crash.

    inst_by_name = {r["institution"]: int(r["code"]) ...} runs before any
    error collection even starts -- a bad code here used to crash the whole
    import with a raw ValueError traceback.
    """
    _patch_dirs(monkeypatch, tmp_path)
    xlsx = _write_workbook(
        tmp_path / "pod.xlsx",
        pod_rows=[],
        project_rows=[],
        revision_rows=[],
        institution_rows=[
            ["XX", "BP Migas", "POD 2003 to 2012"],
            [4, "SKK Migas", "POD 2013 onward"],
        ],
    )
    with pytest.raises(PodRegistryImportError) as exc:
        import_pod_registry_workbook(xlsx)
    assert exc.value.errors == ["institution row 2: invalid code 'XX'"]
