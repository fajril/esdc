from pathlib import Path

import fitz  # PyMuPDF, pulled in by pymupdf4llm
import pytest

from esdc.corpus import extractor
from esdc.corpus.extractor import (
    SUPPORTED_EXTENSIONS,
    ExtractionResult,
    extract_document,
    extract_docx,
    extract_markdown,
    extract_pdf,
)


@pytest.fixture
def text_pdf(tmp_path: Path) -> Path:
    """One page with a real text layer."""
    path = tmp_path / "surat.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Nomor: SRT-001/SKK/2026")
    page.insert_text((72, 100), "Perihal: Persetujuan POD Lapangan Duri tahun 2026")
    doc.save(path)
    return path


@pytest.fixture
def mixed_pdf(tmp_path: Path) -> Path:
    """Page 1 has text, page 2 is blank (scanned-page stand-in)."""
    path = tmp_path / "mixed.pdf"
    doc = fitz.open()
    page = doc.new_page()
    # insert_textbox wraps within the rect; a bare insert_text() line this
    # long overflows the page width and pymupdf4llm drops the whole span.
    page.insert_textbox(
        (72, 72, 523, 200),
        "Halaman digital dengan teks yang cukup panjang " * 3,
    )
    doc.new_page()
    doc.save(path)
    return path


@pytest.fixture
def scanned_pdf(tmp_path: Path) -> Path:
    """One blank page (scanned-page stand-in): no text layer at all."""
    path = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(path)
    return path


class FakeOcr:
    def ocr_page(self, png_bytes: bytes) -> str:
        return "teks hasil OCR halaman scan"


def test_native_only(text_pdf):
    result = extract_pdf(text_pdf, ocr_client=None)
    assert isinstance(result, ExtractionResult)
    assert "SRT-001" in result.markdown
    assert "<!-- page 1: native -->" in result.markdown
    assert result.method == "native"
    assert result.pages_ocr == 0 and result.pages_native == 1


def test_mixed_uses_ocr_per_page(mixed_pdf):
    result = extract_pdf(mixed_pdf, ocr_client=FakeOcr())
    assert result.method == "mixed"
    assert "<!-- page 1: native -->" in result.markdown
    assert "<!-- page 2: llm_ocr -->" in result.markdown
    assert "hasil OCR" in result.markdown
    assert result.pages_native == 1 and result.pages_ocr == 1


def test_all_scanned_uses_llm_ocr(scanned_pdf):
    result = extract_pdf(scanned_pdf, ocr_client=FakeOcr())
    assert result.method == "llm_ocr"
    assert result.pages_native == 0 and result.pages_ocr == 1
    assert "<!-- page 1: llm_ocr -->" in result.markdown


def test_scanned_without_ocr_raises(mixed_pdf):
    with pytest.raises(ValueError, match="scanned page"):
        extract_pdf(mixed_pdf, ocr_client=None)


def _pdf_with_image(tmp_path: Path, name: str, img_rect: fitz.Rect) -> Path:
    """Page with a healthy text layer plus one embedded raster image."""
    path = tmp_path / name
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 300, 200), False)
    pix.clear_with(200)
    doc = fitz.open()
    page = doc.new_page()  # 595 x 842 pt
    page.insert_textbox(
        (72, 72, 523, 200),
        "Halaman digital dengan teks yang cukup panjang " * 3,
    )
    page.insert_image(img_rect, pixmap=pix)
    doc.save(path)
    doc.close()
    return path


@pytest.fixture
def image_table_pdf(tmp_path: Path) -> Path:
    """Big embedded image (~26% of page area): a table saved as a picture."""
    return _pdf_with_image(tmp_path, "table.pdf", fitz.Rect(72, 250, 500, 550))


@pytest.fixture
def logo_pdf(tmp_path: Path) -> Path:
    """Tiny embedded image (~0.5% of page area): a letterhead logo."""
    return _pdf_with_image(tmp_path, "logo.pdf", fitz.Rect(72, 250, 122, 300))


class FakeImageOcr:
    def __init__(self) -> None:
        self.calls = 0

    def ocr_page(self, png_bytes: bytes) -> str:
        self.calls += 1
        return "| Milestone | Plan |\n|---|---|\n| KOM | 6 Mei 2025 |"


def test_native_page_big_image_gets_ocr(image_table_pdf):
    ocr = FakeImageOcr()
    result = extract_pdf(image_table_pdf, ocr_client=ocr)
    assert "<!-- page 1: native -->" in result.markdown
    assert "<!-- page 1 image 1: llm_ocr -->" in result.markdown
    assert "| KOM | 6 Mei 2025 |" in result.markdown
    # native text still present and untouched
    assert "Halaman digital" in result.markdown
    assert ocr.calls == 1
    assert result.images_ocr == 1
    assert result.images_skipped == 0
    assert result.pages_native == 1 and result.pages_ocr == 0
    # LLM content present -> no longer pure native
    assert result.method == "mixed"


def test_small_image_ignored(logo_pdf):
    ocr = FakeImageOcr()
    result = extract_pdf(logo_pdf, ocr_client=ocr)
    assert ocr.calls == 0
    assert result.images_ocr == 0
    assert result.images_skipped == 0
    assert result.method == "native"
    assert "image" not in result.markdown


def test_big_image_without_ocr_client_skipped_not_fatal(image_table_pdf):
    result = extract_pdf(image_table_pdf, ocr_client=None)
    assert result.images_ocr == 0
    assert result.images_skipped == 1
    assert result.method == "native"
    assert "<!-- page 1 image" not in result.markdown
    assert "Halaman digital" in result.markdown


# --------------------------------------------------------------------------
# extract_markdown
# --------------------------------------------------------------------------


def test_extract_markdown_verbatim(tmp_path):
    src = tmp_path / "notes.md"
    src.write_text("# Judul\n\nisi dokumen", encoding="utf-8")
    r = extract_markdown(src)
    assert r.markdown.startswith("<!-- page 1: native_md -->")
    assert "# Judul" in r.markdown and "isi dokumen" in r.markdown
    assert (r.method, r.page_count, r.pages_native, r.pages_ocr) == (
        "native_md",
        1,
        1,
        0,
    )


def test_extract_markdown_empty_raises(tmp_path):
    src = tmp_path / "empty.md"
    src.write_text("   \n", encoding="utf-8")
    with pytest.raises(ValueError, match="no extractable"):
        extract_markdown(src)


# --------------------------------------------------------------------------
# extract_docx
# --------------------------------------------------------------------------


def _tiny_png_bytes() -> bytes:
    """A tiny (4x4) real PNG, generated via fitz so no extra image dep is needed."""
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 4, 4), False)
    pix.clear_with(200)
    return pix.tobytes("png")


def make_docx(
    tmp_path: Path,
    name: str = "doc.docx",
    *,
    headings: tuple[str, ...] = (),
    paragraphs: tuple[str, ...] = (),
    table: list[list[str]] | None = None,
    image: bool = False,
) -> Path:
    from io import BytesIO

    from docx import Document

    d = Document()
    for h in headings:
        d.add_heading(h, level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    if image:
        d.add_picture(BytesIO(_tiny_png_bytes()))
    path = tmp_path / name
    d.save(path)
    return path


def test_extract_docx_headings_paragraphs_in_order(tmp_path):
    path = make_docx(
        tmp_path,
        headings=("Judul",),
        paragraphs=("paragraf pertama", "paragraf kedua"),
    )
    r = extract_docx(path)
    assert "<!-- page 1: native_docx -->" in r.markdown
    assert r.method == "native_docx"
    assert (r.page_count, r.pages_native, r.pages_ocr) == (1, 1, 0)
    idx_heading = r.markdown.index("# Judul")
    idx_p1 = r.markdown.index("paragraf pertama")
    idx_p2 = r.markdown.index("paragraf kedua")
    assert idx_heading < idx_p1 < idx_p2


def test_extract_docx_table_to_markdown_table(tmp_path):
    path = make_docx(
        tmp_path,
        table=[["Nama", "Jabatan"], ["Andi", "Kepala"], ["A|B", "C"]],
    )
    r = extract_docx(path)
    assert "| Nama | Jabatan |" in r.markdown
    assert "| --- | --- |" in r.markdown
    assert "| Andi | Kepala |" in r.markdown
    assert "| A\\|B | C |" in r.markdown


def test_extract_docx_image_counted_skipped(tmp_path):
    path = make_docx(tmp_path, paragraphs=("teks",), image=True)
    r = extract_docx(path)
    assert r.images_skipped == 1
    assert r.images_ocr == 0


def test_extract_docx_empty_raises(tmp_path):
    from docx import Document

    d = Document()
    path = tmp_path / "empty.docx"
    d.save(path)
    with pytest.raises(ValueError, match="no extractable"):
        extract_docx(path)


# --------------------------------------------------------------------------
# extract_document dispatcher
# --------------------------------------------------------------------------


def test_supported_extensions_contains_pdf_docx_md():
    assert SUPPORTED_EXTENSIONS == (".pdf", ".docx", ".md")


def test_extract_document_routes_pdf(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(
        extractor, "extract_pdf", lambda path, ocr_client, a, b, c: calls.append(
            ("pdf", path, ocr_client)
        )
    )
    path = tmp_path / "doc.pdf"
    path.write_bytes(b"%PDF-1.4")
    extract_document(path, ocr_client="sentinel-ocr")
    assert calls == [("pdf", path, "sentinel-ocr")]


def test_extract_document_routes_docx(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(
        extractor, "extract_docx", lambda path: calls.append(("docx", path))
    )
    path = tmp_path / "doc.docx"
    path.write_bytes(b"PK")
    extract_document(path, ocr_client=None)
    assert calls == [("docx", path)]


def test_extract_document_routes_markdown(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(
        extractor, "extract_markdown", lambda path: calls.append(("md", path))
    )
    path = tmp_path / "notes.md"
    path.write_text("hi")
    extract_document(path, ocr_client=None)
    assert calls == [("md", path)]


def test_extract_document_unsupported_extension_raises(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("hi")
    with pytest.raises(ValueError, match="unsupported format"):
        extract_document(path, ocr_client=None)
