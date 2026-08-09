"""Tests for the load CLI command."""

from __future__ import annotations

import duckdb
import pandas as pd
import pytest
import yaml
from openpyxl import load_workbook
from openpyxl.cell import Cell

from esdc.chat.tools import knowledge_traversal
from esdc.configs import Config
from esdc.esdc import app
from esdc.loaders import (
    POD_MONITORING_SHEET_NAME,
    POD_PLAN_SHEET_NAME,
    POD_PROJECT_SHEET_NAME,
    POD_SCHEMA_PATH,
    LoadSchemaError,
    build_loaded_schema_graph,
    load_schema_from_yaml,
)
from esdc.pod_registry.store import get_sqlite_connection


def _write_schema(path, *, columns=None, links=None):
    schema = {
        "table_name": "sample_table",
        "description": "Sample loaded table",
        "sheet_name": "Sheet1",
        "columns": columns
        or [
            {
                "name": "project_id",
                "type": "string",
                "description": "ID proyek",
                "aliases": ["project code", "kode proyek"],
            },
            {
                "name": "volume",
                "type": "double",
                "description": "Volume sumber daya",
                "unit": "MMBOE",
            },
        ],
    }
    if links is not None:
        schema["links"] = links
    path.write_text(yaml.safe_dump(schema, sort_keys=False), encoding="utf-8")
    return path


def _write_excel(path, rows):
    df = pd.DataFrame(rows)
    df.to_excel(path, index=False, sheet_name="Sheet1")
    return path


def _write_pod_excel(path):
    plan_row, project_row, monitoring_row = _make_pod_rows(
        project_id="PROJECT-CANONICAL"
    )
    return _write_pod_workbook(path, [plan_row], [project_row], [monitoring_row])


class TestLoadSchemaParser:
    def test_schema_valid(self, tmp_path):
        schema = load_schema_from_yaml(_write_schema(tmp_path / "schema.yaml"))

        assert schema.table_name == "sample_table"
        assert schema.sheet_name == "Sheet1"
        assert schema.columns[0].name == "project_id"
        assert schema.columns[1].duckdb_type == "DOUBLE"

    def test_pod_schema_template_loads(self):
        schema = load_schema_from_yaml(POD_SCHEMA_PATH)

        assert schema.table_name == "pod_value_case"
        assert schema.sheet_name == "POD"
        assert schema.links == ()
        names = {column.name for column in schema.columns}
        assert {
            "pod_id",
            "case_type",
            "report_date",
            "as_of_date",
            "pod_scope",
        } <= names
        assert (
            not {"pod_name", "pod_letter_num", "supercedes_by", "effective_date"}
            & names
        )
        assert any("ProducingLicense" in column.maps_to for column in schema.columns)

    def test_schema_missing_required_field(self, tmp_path):
        path = tmp_path / "schema.yaml"
        path.write_text("table_name: sample_table\n", encoding="utf-8")

        with pytest.raises(LoadSchemaError, match="description"):
            load_schema_from_yaml(path)

    def test_schema_duplicate_column_name(self, tmp_path):
        path = _write_schema(
            tmp_path / "schema.yaml",
            columns=[
                {"name": "project_id", "type": "string", "description": "A"},
                {"name": "project_id", "type": "string", "description": "B"},
            ],
        )

        with pytest.raises(LoadSchemaError, match="Duplicate column name"):
            load_schema_from_yaml(path)

    def test_schema_unknown_type(self, tmp_path):
        path = _write_schema(
            tmp_path / "schema.yaml",
            columns=[
                {"name": "project_id", "type": "unknown", "description": "A"},
            ],
        )

        with pytest.raises(LoadSchemaError, match="Unsupported column type"):
            load_schema_from_yaml(path)

    def test_schema_valid_with_links_and_maps_to(self, tmp_path):
        path = _write_schema(
            tmp_path / "schema.yaml",
            columns=[
                {"name": "project_id", "type": "string", "description": "Project ID"},
                {
                    "name": "volume",
                    "type": "double",
                    "description": "Volume",
                    "maps_to": ["RESOURCES"],
                },
            ],
            links=[
                {
                    "entity_type": "Project",
                    "column": "project_id",
                }
            ],
        )

        schema = load_schema_from_yaml(path)

        assert schema.links[0].entity_type == "Project"
        assert schema.links[0].target_key == "project_id"
        assert schema.columns[1].maps_to == ("RESOURCES",)

    def test_schema_invalid_link_entity_type(self, tmp_path):
        path = _write_schema(
            tmp_path / "schema.yaml",
            links=[{"entity_type": "Basin", "column": "project_id"}],
        )

        with pytest.raises(LoadSchemaError, match="Unsupported links"):
            load_schema_from_yaml(path)

    def test_schema_link_column_must_exist(self, tmp_path):
        path = _write_schema(
            tmp_path / "schema.yaml",
            links=[{"entity_type": "Project", "column": "missing_id"}],
        )

        with pytest.raises(LoadSchemaError, match="is not in columns"):
            load_schema_from_yaml(path)

    def test_loaded_schema_graph_builds_metadata_edges(self, tmp_path):
        path = _write_schema(
            tmp_path / "schema.yaml",
            columns=[
                {"name": "project_id", "type": "string", "description": "Project ID"},
                {
                    "name": "volume",
                    "type": "double",
                    "description": "Volume",
                    "maps_to": ["RESOURCES"],
                },
            ],
            links=[{"entity_type": "Project", "column": "project_id"}],
        )
        schema = load_schema_from_yaml(path)
        schema_yaml = yaml.safe_dump(schema.raw, sort_keys=False)

        conn = build_loaded_schema_graph(
            [("sample_table", "Sample loaded table", "Sheet1", schema_yaml)]
        )
        result = conn.execute(
            "MATCH (t:LoadedTable)-[:LOADED_TABLE_HAS_COLUMN]->(c:LoadedColumn) "
            "RETURN t.table_name, c.name"
        )
        rows = []
        while result.has_next():
            rows.append(tuple(result.get_next()))

        assert rows == [("sample_table", "project_id"), ("sample_table", "volume")]


class TestLoadCommand:
    def test_load_help(self, runner):
        result = runner.invoke(app, ["load", "--help"])

        assert result.exit_code == 0
        assert "--from-excel" in result.stdout
        assert "--schema" in result.stdout
        assert "--schema-pod" in result.stdout
        assert "esdc schema --generate --from-excel data.xlsx" in result.stdout

    def test_load_requires_schema_or_schema_pod(self, runner, tmp_path):
        excel_path = _write_excel(tmp_path / "data.xlsx", [{"project_id": "A-1"}])

        result = runner.invoke(app, ["load", "--from-excel", str(excel_path)])

        assert result.exit_code == 1
        assert (
            "specify exactly one of --schema, --schema-pod, or --pod-registry"
            in result.stdout
        )

    def test_load_excel_creates_table_and_metadata(
        self, runner, isolated_config, tmp_path
    ):
        schema_path = _write_schema(tmp_path / "schema.yaml")
        excel_path = _write_excel(
            tmp_path / "data.xlsx",
            [
                {"project_id": "A-1", "volume": 12.5},
                {"project_id": "A-2", "volume": 20.0},
            ],
        )

        result = runner.invoke(
            app,
            ["load", "--from-excel", str(excel_path), "--schema", str(schema_path)],
        )

        assert result.exit_code == 0
        assert "Loaded 2 rows x 2 columns into 'sample_table'" in result.stdout

        conn = duckdb.connect(str(Config.get_db_file()))
        try:
            count_row = conn.execute("SELECT COUNT(*) FROM sample_table").fetchone()
            assert count_row is not None
            count = count_row[0]
            metadata = conn.execute(
                "SELECT description FROM _loaded_table_schemas WHERE table_name = ?",
                ["sample_table"],
            ).fetchone()
        finally:
            conn.close()
        assert count == 2
        assert metadata == ("Sample loaded table",)

    def test_load_pod_schema_builtin(self, runner, isolated_config, tmp_path):
        _seed_canonical_pod()
        excel_path = _write_pod_excel(tmp_path / "pod.xlsx")

        result = runner.invoke(
            app,
            ["load", "--from-excel", str(excel_path), "--schema-pod"],
        )

        assert result.exit_code == 0
        assert "into 'pod_plan'" in result.stdout
        assert "into 'pod_monitoring'" in result.stdout
        conn = duckdb.connect(str(Config.get_db_file()))
        try:
            plan_rows = conn.execute(
                "SELECT case_type, pod_id FROM pod_plan"
            ).fetchall()
            monitoring_rows = conn.execute(
                "SELECT case_type, pod_id FROM pod_monitoring"
            ).fetchall()
            eco_meta = conn.execute(
                "SELECT description FROM _loaded_table_schemas"
                " WHERE table_name = 'pod_economics'"
            ).fetchone()
        finally:
            conn.close()
        assert plan_rows == [("plan", "PL-2024-0001-1-1-0")]
        assert monitoring_rows == [("outlook", "PL-2024-0001-1-1-0")]
        assert eco_meta is not None

        sconn = get_sqlite_connection()
        try:
            value_cases = sconn.execute(
                "SELECT pod_id, case_type, report_date, as_of_date FROM pod_value_case"
                " ORDER BY case_type"
            ).fetchall()
        finally:
            sconn.close()
        assert [tuple(row) for row in value_cases] == [
            ("PL-2024-0001-1-1-0", "outlook", "2024-04-01", "2024-03-31"),
            ("PL-2024-0001-1-1-0", "plan", "2024-01-01", "2024-12-31"),
        ]

    def test_load_replaces_existing_table(self, runner, isolated_config, tmp_path):
        schema_path = _write_schema(tmp_path / "schema.yaml")
        first_excel = _write_excel(
            tmp_path / "data1.xlsx",
            [{"project_id": "A-1", "volume": 12.5}],
        )
        second_excel = _write_excel(
            tmp_path / "data2.xlsx",
            [
                {"project_id": "B-1", "volume": 99.0},
                {"project_id": "B-2", "volume": 100.0},
            ],
        )

        first = runner.invoke(
            app,
            ["load", "--from-excel", str(first_excel), "--schema", str(schema_path)],
        )
        second = runner.invoke(
            app,
            ["load", "--from-excel", str(second_excel), "--schema", str(schema_path)],
        )

        assert first.exit_code == 0
        assert second.exit_code == 0
        conn = duckdb.connect(str(Config.get_db_file()))
        try:
            rows = conn.execute(
                "SELECT project_id FROM sample_table ORDER BY project_id"
            ).fetchall()
        finally:
            conn.close()
        assert rows == [("B-1",), ("B-2",)]

    def test_load_excel_column_mismatch_fails(self, runner, isolated_config, tmp_path):
        schema_path = _write_schema(tmp_path / "schema.yaml")
        excel_path = _write_excel(
            tmp_path / "data.xlsx",
            [{"project_id": "A-1", "unexpected": 12.5}],
        )

        result = runner.invoke(
            app,
            ["load", "--from-excel", str(excel_path), "--schema", str(schema_path)],
        )

        assert result.exit_code == 1
        assert "Excel columns do not match schema" in result.stdout
        assert "missing columns: volume" in result.stdout
        assert "unexpected columns: unexpected" in result.stdout

    def test_load_excel_missing_sheet_fails(self, runner, isolated_config, tmp_path):
        schema_path = _write_schema(tmp_path / "schema.yaml")
        excel_path = tmp_path / "data.xlsx"
        pd.DataFrame([{"project_id": "A-1", "volume": 12.5}]).to_excel(
            excel_path, index=False, sheet_name="OtherSheet"
        )

        result = runner.invoke(
            app,
            ["load", "--from-excel", str(excel_path), "--schema", str(schema_path)],
        )

        assert result.exit_code == 1
        assert "Excel sheet 'Sheet1' not found" in result.stdout

    def test_load_excel_type_conversion_fails(self, runner, isolated_config, tmp_path):
        schema_path = _write_schema(tmp_path / "schema.yaml")
        excel_path = _write_excel(
            tmp_path / "data.xlsx",
            [{"project_id": "A-1", "volume": "not-a-number"}],
        )

        result = runner.invoke(
            app,
            ["load", "--from-excel", str(excel_path), "--schema", str(schema_path)],
        )

        assert result.exit_code == 1
        assert "Column 'volume' cannot be converted to double" in result.stdout

    def test_load_with_links_warns_for_unmatched_ids(
        self, runner, isolated_config, tmp_path
    ):
        Config.init_config()
        Config.get_db_dir().mkdir(parents=True, exist_ok=True)
        conn = duckdb.connect(str(Config.get_db_file()))
        try:
            conn.execute("CREATE TABLE project_resources (project_id TEXT)")
            conn.execute("INSERT INTO project_resources VALUES ('A-1')")
        finally:
            conn.close()
        schema_path = _write_schema(
            tmp_path / "schema.yaml",
            links=[{"entity_type": "Project", "column": "project_id"}],
        )
        excel_path = _write_excel(
            tmp_path / "data.xlsx",
            [
                {"project_id": "A-1", "volume": 12.5},
                {"project_id": "A-2", "volume": 20.0},
            ],
        )

        result = runner.invoke(
            app,
            ["load", "--from-excel", str(excel_path), "--schema", str(schema_path)],
        )

        assert result.exit_code == 0
        assert (
            "Warning: 1 value(s) in 'project_id' did not match Project.project_id"
            in result.stdout
        )

    def test_knowledge_traversal_finds_loaded_table_and_column(
        self, runner, isolated_config, tmp_path
    ):
        schema_path = _write_schema(
            tmp_path / "schema.yaml",
            columns=[
                {
                    "name": "project_id",
                    "type": "string",
                    "description": "ID proyek",
                    "aliases": ["project code", "kode proyek"],
                },
                {
                    "name": "volume",
                    "type": "double",
                    "description": "Volume sumber daya",
                    "unit": "MMBOE",
                    "maps_to": ["RESOURCES"],
                },
            ],
            links=[{"entity_type": "Project", "column": "project_id"}],
        )
        excel_path = _write_excel(
            tmp_path / "data.xlsx",
            [{"project_id": "A-1", "volume": 12.5}],
        )
        load_result = runner.invoke(
            app,
            ["load", "--from-excel", str(excel_path), "--schema", str(schema_path)],
        )
        assert load_result.exit_code == 0

        table_result = knowledge_traversal.invoke(
            {"topic": "definition", "entity": "sample_table"}
        )
        column_result = knowledge_traversal.invoke(
            {"topic": "definition", "entity": "kode proyek"}
        )

        assert "Loaded table: sample_table" in table_result
        assert "project_id: ID proyek" in table_result
        assert "project_id links to Project.project_id" in table_result
        assert "Loaded column: sample_table.project_id" in column_result
        assert "Description: ID proyek" in column_result
        volume_result = knowledge_traversal.invoke(
            {"topic": "definition", "entity": "volume"}
        )
        assert "Maps to KSMI: RESOURCES" in volume_result
        assert "A-1" not in table_result
        assert "A-1" not in column_result


def _seed_canonical_pod(project_id: str = "PROJECT-CANONICAL"):
    """Seed the canonical SQLite POD registry row the POD tests load against."""
    conn = get_sqlite_connection()
    conn.execute(
        "INSERT INTO r_institution (code, institution) VALUES (1, 'SKK Migas')"
    )
    conn.execute("INSERT INTO r_pod_type (code, pod_type) VALUES (1, 'POD I')")
    conn.execute(
        "INSERT INTO m_pod (id, pod_id, pod_name, pod_letter_num, approval_date,"
        " institution_code, pod_type_code, rev_num, approval_seq)"
        " VALUES (1, 'PL-2024-0001-1-1-0', 'POD Canonical', 'L-CANONICAL',"
        " '2024-01-15', 1, 1, 0, 1)"
    )
    conn.execute(
        "INSERT INTO project_pod (pod_id, project_id) VALUES (1, ?)", [project_id]
    )
    conn.commit()
    conn.close()


def _make_pod_rows(
    plan_pod_id: str = "PL-2024-0001-1-1-0",
    project_pod_id: str = "PL-2024-0001-1-1-0",
    monitoring_pod_id: str = "PL-2024-0001-1-1-0",
    project_id: str = "PROJECT-CANONICAL",
):
    schema = load_schema_from_yaml(POD_SCHEMA_PATH)
    metric_columns = [
        column
        for column in schema.columns
        if column.name
        not in {"report_date", "as_of_date", "case_type", "pod_id", "pod_scope"}
    ]
    plan_row = {
        "pod_id": plan_pod_id,
        "pod_scope": "Development",
        "report_date": pd.Timestamp("2024-01-01"),
        "as_of_date": pd.Timestamp("2024-12-31"),
    }
    monitoring_row = {
        "pod_id": monitoring_pod_id,
        "case_type": "outlook",
        "pod_scope": "Development",
        "report_date": pd.Timestamp("2024-04-01"),
        "as_of_date": pd.Timestamp("2024-03-31"),
    }
    for column in metric_columns:
        plan_row[column.name] = 1.0
        monitoring_row[column.name] = 2.0
    project_row = {"pod_id": project_pod_id, "project_id": project_id}
    return plan_row, project_row, monitoring_row


def _write_pod_workbook(path, plan_rows, project_rows, monitoring_rows):
    """Write a POD workbook from row dicts; project_rows=None omits the sheet."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(plan_rows).to_excel(
            writer, index=False, sheet_name=POD_PLAN_SHEET_NAME, startrow=1
        )
        if project_rows is not None:
            pd.DataFrame(project_rows).to_excel(
                writer, index=False, sheet_name=POD_PROJECT_SHEET_NAME, startrow=1
            )
        pd.DataFrame(monitoring_rows).to_excel(
            writer, index=False, sheet_name=POD_MONITORING_SHEET_NAME, startrow=1
        )
    workbook = load_workbook(path)
    try:
        descriptions = {
            POD_PLAN_SHEET_NAME: "Baseline POD plan data.",
            POD_PROJECT_SHEET_NAME: "Many-to-many POD to project mapping.",
            POD_MONITORING_SHEET_NAME: "POD outlook and actual monitoring data.",
        }
        for sheet_name, description in descriptions.items():
            if sheet_name not in workbook.sheetnames:
                continue
            ws = workbook[sheet_name]
            ws.merge_cells(
                start_row=1, start_column=1, end_row=1, end_column=ws.max_column
            )
            title_cell = ws.cell(1, 1)
            assert isinstance(title_cell, Cell)
            title_cell.value = description
    finally:
        workbook.save(path)
        workbook.close()
    return path


def _write_pod_excel_with_data(
    path, plan_row, project_row, monitoring_row, *, include_project=True
):
    return _write_pod_workbook(
        path,
        [plan_row],
        [project_row] if include_project else None,
        [monitoring_row],
    )


class TestPodLoadValidation:
    def test_load_pod_views_exist(self, runner, isolated_config, tmp_path):
        _seed_canonical_pod()
        excel_path = _write_pod_excel(tmp_path / "pod.xlsx")
        result = runner.invoke(
            app, ["load", "--from-excel", str(excel_path), "--schema-pod"]
        )
        assert result.exit_code == 0

        conn = duckdb.connect(str(Config.get_db_file()))
        try:
            views = conn.execute(
                "SELECT table_name FROM information_schema.tables WHERE table_type = 'VIEW'"
            ).fetchall()
            view_names = {v[0] for v in views}
            assert "pod_economics" in view_names
            assert "pod_plan" in view_names
            assert "pod_monitoring" in view_names
            assert "pod_project_economics" in view_names

            eco_rows = conn.execute(
                "SELECT case_type, pod_id, lifting_oil FROM pod_economics ORDER BY case_type"
            ).fetchall()
            assert len(eco_rows) == 2
            assert eco_rows[0][0] == "outlook"
            assert eco_rows[0][2] == 2.0
            assert eco_rows[1][0] == "plan"
            assert eco_rows[1][2] == 1.0

            plan_scope_row = conn.execute(
                "SELECT pod_scope FROM pod_economics WHERE case_type = 'plan'"
            ).fetchone()
            mon_scope_row = conn.execute(
                "SELECT pod_scope FROM pod_economics WHERE case_type = 'outlook'"
            ).fetchone()
            assert plan_scope_row is not None
            assert mon_scope_row is not None
            assert plan_scope_row[0] == "Development"
            assert mon_scope_row[0] == "Development"

            proj_eco = conn.execute(
                "SELECT pod_id, project_id, case_type FROM pod_project_economics ORDER BY case_type"
            ).fetchall()
            assert len(proj_eco) == 2
            assert proj_eco[0][1] == "PROJECT-CANONICAL"
            assert proj_eco[1][1] == "PROJECT-CANONICAL"
        finally:
            conn.close()

    def test_load_pod_plan_same_pod_id_distinct_as_of_dates_stores_both(
        self, runner, isolated_config, tmp_path
    ):
        """Duplicate pod_id is legal when plan as_of_date differs."""
        _seed_canonical_pod()
        plan_row, project_row, monitoring_row = _make_pod_rows(
            project_id="PROJECT-CANONICAL"
        )
        second_plan_row = dict(plan_row)
        second_plan_row["as_of_date"] = pd.Timestamp("2025-12-31")
        excel_path = _write_pod_workbook(
            tmp_path / "pod.xlsx",
            [plan_row, second_plan_row],
            [project_row],
            [monitoring_row],
        )
        result = runner.invoke(
            app, ["load", "--from-excel", str(excel_path), "--schema-pod"]
        )
        assert result.exit_code == 0

        sconn = get_sqlite_connection()
        try:
            rows = sconn.execute(
                "SELECT pod_id, case_type, as_of_date FROM pod_value_case"
                " WHERE case_type = 'plan' ORDER BY as_of_date"
            ).fetchall()
        finally:
            sconn.close()
        assert [tuple(row) for row in rows] == [
            ("PL-2024-0001-1-1-0", "plan", "2024-12-31"),
            ("PL-2024-0001-1-1-0", "plan", "2025-12-31"),
        ]

    def test_load_pod_broken_project_pod_id_fails(
        self, runner, isolated_config, tmp_path
    ):
        _seed_canonical_pod()
        plan_row, _, monitoring_row = _make_pod_rows()
        project_row = {"pod_id": "PL-2024-0002-1-1-1", "project_id": "PRJ-001"}
        excel_path = _write_pod_excel_with_data(
            tmp_path / "pod.xlsx", plan_row, project_row, monitoring_row
        )
        result = runner.invoke(
            app, ["load", "--from-excel", str(excel_path), "--schema-pod"]
        )
        assert result.exit_code == 1
        assert (
            "pod_project references pod_id values not found in pod_plan"
            in result.stdout
        )

    def test_load_pod_broken_monitoring_pod_id_fails(
        self, runner, isolated_config, tmp_path
    ):
        _seed_canonical_pod()
        plan_row, project_row, _ = _make_pod_rows()
        monitoring_row = {
            "pod_id": "PL-2024-0002-1-1-1",
            "case_type": "outlook",
            "pod_scope": "Development",
            "report_date": pd.Timestamp("2024-04-01"),
            "as_of_date": pd.Timestamp("2024-03-31"),
        }
        schema = load_schema_from_yaml(POD_SCHEMA_PATH)
        metric_columns = [
            col
            for col in schema.columns
            if col.name
            not in {"report_date", "as_of_date", "case_type", "pod_id", "pod_scope"}
        ]
        for col in metric_columns:
            monitoring_row[col.name] = 2.0
        excel_path = _write_pod_excel_with_data(
            tmp_path / "pod.xlsx", plan_row, project_row, monitoring_row
        )
        result = runner.invoke(
            app, ["load", "--from-excel", str(excel_path), "--schema-pod"]
        )
        assert result.exit_code == 1
        assert (
            "pod_monitoring references pod_id values not found in pod_plan"
            in result.stdout
        )

    def test_load_pod_unknown_pod_id_fails_before_committing_any_value_case(
        self, runner, isolated_config, tmp_path
    ):
        """A plan row with a pod_id absent from the canonical SQLite registry fails.

        Validation happens against SQLite m_pod before any truth row is
        replaced, so the load exits non-zero and pod_value_case keeps the
        pre-existing plan committed by the earlier good load.
        """
        _seed_canonical_pod()
        good_path = _write_pod_excel(tmp_path / "good.xlsx")
        good = runner.invoke(
            app, ["load", "--from-excel", str(good_path), "--schema-pod"]
        )
        assert good.exit_code == 0

        plan_row, project_row, monitoring_row = _make_pod_rows(
            plan_pod_id="PL-2024-0009-1-1-0",
            project_pod_id="PL-2024-0009-1-1-0",
            monitoring_pod_id="PL-2024-0009-1-1-0",
        )
        bad_path = _write_pod_excel_with_data(
            tmp_path / "bad.xlsx",
            plan_row,
            project_row,
            monitoring_row,
            include_project=False,
        )
        bad = runner.invoke(
            app, ["load", "--from-excel", str(bad_path), "--schema-pod"]
        )
        assert bad.exit_code == 1
        assert "unknown pod_id" in bad.stdout

        sconn = get_sqlite_connection()
        try:
            rows = sconn.execute(
                "SELECT pod_id, case_type FROM pod_value_case ORDER BY case_type"
            ).fetchall()
        finally:
            sconn.close()
        assert [tuple(row) for row in rows] == [
            ("PL-2024-0001-1-1-0", "outlook"),
            ("PL-2024-0001-1-1-0", "plan"),
        ]

    def test_load_pod_workbook_project_mismatch_fails_without_committing(
        self, runner, isolated_config, tmp_path
    ):
        """A pod_project sheet contradicting the canonical SQLite mapping fails.

        The canonical project_pod is the source of truth for pod_project.
        Loading an analytical workbook that maps the canonical pod
        PL-2024-0001-1-1-0 -> PROJECT-WORKBOOK (canonical says
        PROJECT-CANONICAL) must exit non-zero and commit zero value cases.
        """
        _seed_canonical_pod(project_id="PROJECT-CANONICAL")

        plan_row, project_row, monitoring_row = _make_pod_rows(
            project_id="PROJECT-WORKBOOK",
        )
        workbook = _write_pod_excel_with_data(
            tmp_path / "analytical.xlsx", plan_row, project_row, monitoring_row
        )
        result = runner.invoke(
            app, ["load", "--from-excel", str(workbook), "--schema-pod"]
        )
        assert result.exit_code == 1
        assert "does not match canonical project_pod" in result.stdout

        sconn = get_sqlite_connection()
        try:
            exists = sconn.execute(
                "SELECT COUNT(*) FROM sqlite_master"
                " WHERE type='table' AND name='pod_value_case'"
            ).fetchone()
        finally:
            sconn.close()
        assert exists is not None and exists[0] == 0

    def test_load_pod_workbook_project_matching_canonical_loads(
        self, runner, isolated_config, tmp_path
    ):
        """A pod_project sheet matching the canonical SQLite mapping loads."""
        _seed_canonical_pod(project_id="PROJECT-CANONICAL")
        plan_row, project_row, monitoring_row = _make_pod_rows(
            project_id="PROJECT-CANONICAL"
        )
        workbook = _write_pod_excel_with_data(
            tmp_path / "pod.xlsx", plan_row, project_row, monitoring_row
        )
        result = runner.invoke(
            app, ["load", "--from-excel", str(workbook), "--schema-pod"]
        )
        assert result.exit_code == 0

        sconn = get_sqlite_connection()
        try:
            count_row = sconn.execute("SELECT COUNT(*) FROM pod_value_case").fetchone()
        finally:
            sconn.close()
        assert count_row is not None and count_row[0] == 2

    def test_load_pod_projection_failure_reports_committed_sqlite_and_retry(
        self, runner, isolated_config, tmp_path, monkeypatch
    ):
        """If the DuckDB projection fails after commit, report truth committed.

        The CLI must exit non-zero, say the value cases were committed to
        SQLite, and point at `esdc corpus sync` to retry the projection.
        """
        _seed_canonical_pod()
        excel_path = _write_pod_excel(tmp_path / "pod.xlsx")

        def boom():
            raise RuntimeError("projection broke")

        monkeypatch.setattr("esdc.pod_registry.publish.publish_pod_registry", boom)

        result = runner.invoke(
            app, ["load", "--from-excel", str(excel_path), "--schema-pod"]
        )

        assert result.exit_code == 1
        assert "committed to SQLite" in result.stdout
        assert "esdc corpus sync" in result.stdout

        sconn = get_sqlite_connection()
        try:
            count_row = sconn.execute("SELECT COUNT(*) FROM pod_value_case").fetchone()
        finally:
            sconn.close()
        assert count_row is not None
        assert count_row[0] == 2


class TestSchemaCommand:
    def test_schema_help(self, runner):
        result = runner.invoke(app, ["schema", "--help"])

        assert result.exit_code == 0
        assert "--generate" in result.stdout
        assert "--from-excel" in result.stdout
        assert "--output" in result.stdout
        assert "--overwrite" in result.stdout
        assert "--schema-pod" in result.stdout

    def test_schema_pod_copies_template_to_current_directory(
        self, runner, tmp_path, monkeypatch
    ):
        monkeypatch.chdir(tmp_path)

        result = runner.invoke(app, ["schema", "--schema-pod"])

        output_path = tmp_path / "pod_template.xlsx"
        assert result.exit_code == 0
        assert output_path.exists()
        assert f"Generated POD workbook template at {output_path}" in result.stdout
        workbook = load_workbook(output_path)
        try:
            assert workbook.sheetnames == [
                POD_PLAN_SHEET_NAME,
                POD_MONITORING_SHEET_NAME,
            ]
            plan_headers = [cell.value for cell in workbook[POD_PLAN_SHEET_NAME][2]]
            monitoring_headers = [
                cell.value for cell in workbook[POD_MONITORING_SHEET_NAME][2]
            ]
        finally:
            workbook.close()
        assert plan_headers[:4] == ["pod_id", "pod_scope", "report_date", "as_of_date"]
        assert monitoring_headers[:5] == [
            "pod_id",
            "case_type",
            "pod_scope",
            "report_date",
            "as_of_date",
        ]

        assert "pod_plan_ids" in workbook.defined_names
        plan_sheet = workbook[POD_PLAN_SHEET_NAME]
        monitoring_sheet = workbook[POD_MONITORING_SHEET_NAME]
        assert len(plan_sheet.data_validations.dataValidation) == 0
        assert len(monitoring_sheet.data_validations.dataValidation) == 2
        mon_dv_case = monitoring_sheet.data_validations.dataValidation[0]
        mon_dv_pod = monitoring_sheet.data_validations.dataValidation[1]
        assert "pod_plan_ids" in mon_dv_pod.formula1
        assert mon_dv_pod.sqref == "A3:A1048576"
        assert "outlook" in mon_dv_case.formula1
        assert mon_dv_case.sqref == "B3:B1048576"

    def test_schema_pod_existing_output_fails_without_overwrite(self, runner, tmp_path):
        output_path = tmp_path / "pod_template.xlsx"
        output_path.write_text("keep: true\n", encoding="utf-8")

        result = runner.invoke(
            app,
            ["schema", "--schema-pod", "--output", str(output_path)],
        )

        assert result.exit_code == 1
        assert "Output schema already exists" in result.stdout
        assert output_path.read_text(encoding="utf-8") == "keep: true\n"

    def test_generate_schema_default_output_in_current_directory(
        self, runner, tmp_path, monkeypatch
    ):
        monkeypatch.chdir(tmp_path)
        excel_path = _write_excel(
            tmp_path / "Data Source.xlsx",
            [{"project_id": "A-1", "volume": 12.5}],
        )

        result = runner.invoke(
            app,
            ["schema", "--generate", "--from-excel", str(excel_path)],
        )

        output_path = tmp_path / "Data Source.schema.yaml"
        assert result.exit_code == 0
        assert output_path.exists()
        assert f"Schema: {output_path}" in result.stdout
        schema = load_schema_from_yaml(output_path)
        assert schema.table_name == "data_source"
        assert schema.sheet_name == "Sheet1"
        assert [column.name for column in schema.columns] == ["project_id", "volume"]
        assert schema.links[0].entity_type == "Project"
        assert schema.links[0].column == "project_id"

    def test_generate_schema_explicit_output(self, runner, tmp_path):
        excel_path = _write_excel(
            tmp_path / "data.xlsx",
            [{"project_id": "A-1", "volume": 12.5}],
        )
        output_path = tmp_path / "schemas" / "data.yaml"

        result = runner.invoke(
            app,
            [
                "schema",
                "--generate",
                "--from-excel",
                str(excel_path),
                "--output",
                str(output_path),
            ],
        )

        assert result.exit_code == 0
        assert output_path.exists()
        schema = load_schema_from_yaml(output_path)
        assert schema.table_name == "data"
        assert schema.columns[1].type == "double"

    def test_generate_schema_detects_standard_links(self, runner, tmp_path):
        excel_path = _write_excel(
            tmp_path / "linked.xlsx",
            [
                {
                    "project_id": "P-1",
                    "field_id": "F-1",
                    "wk_id": "WK-1",
                    "volume": 12.5,
                }
            ],
        )
        output_path = tmp_path / "linked.schema.yaml"

        result = runner.invoke(
            app,
            [
                "schema",
                "--generate",
                "--from-excel",
                str(excel_path),
                "--output",
                str(output_path),
            ],
        )

        assert result.exit_code == 0
        data = yaml.safe_load(output_path.read_text(encoding="utf-8"))
        assert data["links"] == [
            {
                "entity_type": "Project",
                "column": "project_id",
                "target_key": "project_id",
            },
            {"entity_type": "Field", "column": "field_id", "target_key": "field_id"},
            {
                "entity_type": "WorkingArea",
                "column": "wk_id",
                "target_key": "wk_id",
            },
        ]

    def test_generate_schema_omits_links_without_standard_columns(
        self, runner, tmp_path
    ):
        excel_path = _write_excel(
            tmp_path / "plain.xlsx",
            [{"name": "Alpha", "volume": 12.5}],
        )
        output_path = tmp_path / "plain.schema.yaml"

        result = runner.invoke(
            app,
            [
                "schema",
                "--generate",
                "--from-excel",
                str(excel_path),
                "--output",
                str(output_path),
            ],
        )

        assert result.exit_code == 0
        data = yaml.safe_load(output_path.read_text(encoding="utf-8"))
        assert "links" not in data

    def test_generate_schema_existing_output_fails_without_overwrite(
        self, runner, tmp_path
    ):
        excel_path = _write_excel(
            tmp_path / "data.xlsx",
            [{"project_id": "A-1", "volume": 12.5}],
        )
        output_path = tmp_path / "schema.yaml"
        output_path.write_text("keep: true\n", encoding="utf-8")

        result = runner.invoke(
            app,
            [
                "schema",
                "--generate",
                "--from-excel",
                str(excel_path),
                "--output",
                str(output_path),
            ],
        )

        assert result.exit_code == 1
        assert "Output schema already exists" in result.stdout
        assert output_path.read_text(encoding="utf-8") == "keep: true\n"

    def test_generate_schema_existing_output_overwrites(self, runner, tmp_path):
        excel_path = _write_excel(
            tmp_path / "data.xlsx",
            [{"project_id": "A-1", "volume": 12.5}],
        )
        output_path = tmp_path / "schema.yaml"
        output_path.write_text("old: true\n", encoding="utf-8")

        result = runner.invoke(
            app,
            [
                "schema",
                "--generate",
                "--from-excel",
                str(excel_path),
                "--output",
                str(output_path),
                "--overwrite",
            ],
        )

        assert result.exit_code == 0
        schema = load_schema_from_yaml(output_path)
        assert schema.table_name == "data"
        assert schema.columns[0].description == "TODO: describe project_id"

    def test_generate_schema_infers_supported_types(self, runner, tmp_path):
        excel_path = tmp_path / "typed.xlsx"
        df = pd.DataFrame(
            {
                "name": ["Alpha", "Beta"],
                "count": [1, 2],
                "volume": [12.5, 20.25],
                "active": [True, False],
                "report_date": [
                    pd.Timestamp("2024-01-01"),
                    pd.Timestamp("2024-01-02"),
                ],
                "updated_at": [
                    pd.Timestamp("2024-01-01 12:30"),
                    pd.Timestamp("2024-01-02 13:45"),
                ],
            }
        )
        df.to_excel(excel_path, index=False, sheet_name="Sheet1")
        output_path = tmp_path / "typed.schema.yaml"

        result = runner.invoke(
            app,
            [
                "schema",
                "--generate",
                "--from-excel",
                str(excel_path),
                "--output",
                str(output_path),
            ],
        )

        assert result.exit_code == 0
        schema = load_schema_from_yaml(output_path)
        types = {column.name: column.type for column in schema.columns}
        assert types == {
            "name": "string",
            "count": "integer",
            "volume": "double",
            "active": "boolean",
            "report_date": "date",
            "updated_at": "datetime",
        }


class TestLoadPodRegistryCommand:
    def test_load_pod_registry_invokes_importer(self, runner, monkeypatch, tmp_path):
        called = {}

        def fake_import(path):
            called["path"] = path
            return {"m_pod": 2}

        monkeypatch.setattr("esdc.esdc.import_pod_registry_workbook", fake_import)
        xlsx = tmp_path / "pod.xlsx"
        xlsx.write_bytes(b"")

        result = runner.invoke(
            app, ["load", "--from-excel", str(xlsx), "--pod-registry"]
        )

        assert result.exit_code == 0
        assert called["path"] == xlsx
        assert "m_pod" in result.stdout

    def test_load_requires_exactly_one_mode(self, runner, tmp_path):
        xlsx = tmp_path / "pod.xlsx"
        xlsx.write_bytes(b"")

        result = runner.invoke(
            app, ["load", "--from-excel", str(xlsx), "--pod-registry", "--schema-pod"]
        )

        assert result.exit_code == 1
