# esdc/pod_registry/importer.py
"""One-time seed of the POD registry SQLite db from pod-itb-skk.xlsx."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl

from esdc.pod_registry.publish import publish_pod_registry
from esdc.pod_registry.store import get_sqlite_connection

_REQUIRED_SHEETS = (
    "POD Record", "project_pod", "pod_revision", "institution", "pod_type",
)


class PodRegistryImportError(Exception):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("; ".join(errors))


def _norm(header: object) -> str:
    return str(header).strip().lower().replace(" ", "").replace("__", "_")


def _rows(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    headers = [_norm(h) for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if all(v is None for v in raw):
            continue
        out.append(dict(zip(headers, raw, strict=False)))
    return out


def _iso(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)[:10] if value else None


def _int_cell(
    value: object, sheet: str, row_idx: int, column: str, errors: list[str]
) -> int | None:
    """Parse a workbook cell as int, or record a row error and return None.

    A blank or non-numeric cell raises TypeError/ValueError from a bare
    int() call, which would otherwise crash the whole import with a raw
    traceback instead of surfacing as one more entry in the per-row error
    list that becomes PodRegistryImportError.
    """
    try:
        return int(value)
    except (TypeError, ValueError):
        errors.append(f"{sheet} row {row_idx}: invalid {column} '{value}'")
        return None


def import_pod_registry_workbook(
    xlsx_path: Path | str, sqlite_path: Path | None = None
) -> dict[str, int]:
    """Replace registry contents from the workbook, then publish to DuckDB."""
    xlsx = Path(xlsx_path)
    if not xlsx.exists():
        raise PodRegistryImportError([f"Excel file not found: {xlsx}"])
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    missing = [s for s in _REQUIRED_SHEETS if s not in wb.sheetnames]
    if missing:
        raise PodRegistryImportError([f"missing sheet: {s}" for s in missing])

    institutions = _rows(wb["institution"])
    pod_types = _rows(wb["pod_type"])
    pods = _rows(wb["POD Record"])
    projects = _rows(wb["project_pod"])
    revisions = _rows(wb["pod_revision"])

    errors: list[str] = []

    inst_by_name: dict[str, int] = {}
    institution_rows: list[tuple] = []
    for i, r in enumerate(institutions, start=2):
        code = _int_cell(r.get("code"), "institution", i, "code", errors)
        if code is None:
            continue
        inst_by_name[r["institution"]] = code
        institution_rows.append((code, r["institution"], r.get("description")))

    type_by_name: dict[str, int] = {}
    pod_type_rows: list[tuple] = []
    for i, r in enumerate(pod_types, start=2):
        code = _int_cell(r.get("code"), "pod_type", i, "code", errors)
        if code is None:
            continue
        type_by_name[r["podtype"]] = code
        pod_type_rows.append((code, r["podtype"], r.get("description")))

    m_pod_rows: list[tuple] = []
    seen_pod_ids: set[str] = set()
    seen_seqs: set[int] = set()
    seen_ids: set[int] = set()
    for i, r in enumerate(pods, start=2):
        inst = inst_by_name.get(r.get("institution"))
        ptype = type_by_name.get(r.get("pod_type"))
        pod_id = r.get("pod_id_skk")
        pid = r.get("pod_id_itb")
        seq = r.get("approval_seq")
        if inst is None:
            errors.append(
                f"POD Record row {i}: unknown institution '{r.get('institution')}'"
            )
        if ptype is None:
            errors.append(
                f"POD Record row {i}: unknown pod_type '{r.get('pod_type')}'"
            )
        if pod_id in seen_pod_ids:
            errors.append(f"POD Record row {i}: duplicate pod_id_skk '{pod_id}'")
        if seq in seen_seqs:
            errors.append(f"POD Record row {i}: duplicate approval_seq '{seq}'")
        if pid in seen_ids:
            errors.append(f"POD Record row {i}: duplicate pod_id_itb '{pid}'")
        if pid is None:
            errors.append(f"POD Record row {i}: missing pod_id_itb")
        if seq is None:
            errors.append(f"POD Record row {i}: missing approval_seq")
        if pod_id is None:
            errors.append(f"POD Record row {i}: missing pod_id_skk")
        seen_pod_ids.add(pod_id)
        seen_seqs.add(seq)
        seen_ids.add(pid)
        if inst is None or ptype is None or pid is None or seq is None or pod_id is None:
            continue
        int_pid = _int_cell(pid, "POD Record", i, "pod_id_itb", errors)
        int_seq = _int_cell(seq, "POD Record", i, "approval_seq", errors)
        int_rev = _int_cell(r.get("rev_num") or 0, "POD Record", i, "rev_num", errors)
        if int_pid is None or int_seq is None or int_rev is None:
            continue
        m_pod_rows.append((
            int_pid, str(pod_id), r.get("pod_name"), r.get("pod_letter_num"),
            _iso(r.get("approval_date")), inst, ptype,
            int_rev, int_seq,
        ))

    valid_ids = {row[0] for row in m_pod_rows}
    valid_pod_ids = {row[1] for row in m_pod_rows}
    project_rows: list[tuple] = []
    for i, r in enumerate(projects, start=2):
        pod_id_int = _int_cell(r.get("pod_id"), "project_pod", i, "pod_id", errors)
        if pod_id_int is None:
            continue
        if pod_id_int not in valid_ids:
            errors.append(
                f"project_pod row {i}: pod_id {r['pod_id']} not in POD Record"
            )
            continue
        project_rows.append((pod_id_int, str(r["project_id"])))

    revision_rows: list[tuple] = []
    seen_rev_pairs: set[tuple] = set()
    for i, r in enumerate(revisions, start=2):
        succ, pred = r.get("successor_id"), r.get("predecessor_id")
        if succ not in valid_pod_ids or pred not in valid_pod_ids:
            errors.append(f"pod_revision row {i}: unknown pod_id ({succ}, {pred})")
            continue
        if succ == pred:
            errors.append(f"pod_revision row {i}: self-referencing pod_id ({succ})")
            continue
        if (succ, pred) in seen_rev_pairs:
            errors.append(f"pod_revision row {i}: duplicate pair ({succ}, {pred})")
            continue
        seen_rev_pairs.add((succ, pred))
        revision_rows.append((succ, pred))

    if errors:
        raise PodRegistryImportError(errors)

    conn = get_sqlite_connection(sqlite_path)
    try:
        with conn:  # one transaction
            for table in (
                "pod_revision", "project_pod", "m_pod", "r_pod_type", "r_institution",
            ):
                conn.execute(f"DELETE FROM {table}")
            conn.executemany(
                "INSERT INTO r_institution (code, institution, description)"
                " VALUES (?,?,?)",
                institution_rows,
            )
            conn.executemany(
                "INSERT INTO r_pod_type (code, pod_type, description) VALUES (?,?,?)",
                pod_type_rows,
            )
            conn.executemany(
                "INSERT INTO m_pod"
                " (id, pod_id, pod_name, pod_letter_num, approval_date,"
                " institution_code, pod_type_code, rev_num, approval_seq)"
                " VALUES (?,?,?,?,?,?,?,?,?)",
                m_pod_rows,
            )
            conn.executemany(
                "INSERT INTO project_pod (pod_id, project_id) VALUES (?,?)",
                project_rows,
            )
            conn.executemany(
                "INSERT INTO pod_revision (successor_id, predecessor_id) VALUES (?,?)",
                revision_rows,
            )
        counts = {
            "r_institution": len(institution_rows),
            "r_pod_type": len(pod_type_rows),
            "m_pod": len(m_pod_rows),
            "project_pod": len(project_rows),
            "pod_revision": len(revision_rows),
        }
    finally:
        conn.close()

    publish_pod_registry(sqlite_path=sqlite_path)
    return counts
