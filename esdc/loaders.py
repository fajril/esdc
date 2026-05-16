from __future__ import annotations

import contextlib
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, NamedTuple

import duckdb
import pandas as pd
import typer
import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from esdc.configs import Config
from esdc.dbmanager import _ensure_duckdb_database, get_duckdb_connection

_IDENTIFIER_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_METADATA_TABLE = "_loaded_table_schemas"
_DOMAIN_KNOWLEDGE_DIR = Path(__file__).parent / "chat" / "domain_knowledge"
POD_SCHEMA_PATH = _DOMAIN_KNOWLEDGE_DIR / "pod_schema.yaml"
POD_TEMPLATE_FILENAME = "pod_template.xlsx"
POD_PLAN_SHEET_NAME = "pod_plan"
POD_PROJECT_SHEET_NAME = "pod_project"
POD_MONITORING_SHEET_NAME = "pod_monitoring"
_LINK_TARGET_KEYS = {
    "Project": "project_id",
    "Field": "field_id",
    "WorkingArea": "wk_id",
}
_LINK_REFERENCE_TABLES = ("project_resources", "project_timeseries")

_TYPE_MAP = {
    "string": "TEXT",
    "text": "TEXT",
    "integer": "BIGINT",
    "int": "BIGINT",
    "bigint": "BIGINT",
    "double": "DOUBLE",
    "float": "DOUBLE",
    "number": "DOUBLE",
    "boolean": "BOOLEAN",
    "bool": "BOOLEAN",
    "date": "DATE",
    "datetime": "TIMESTAMP",
}


@dataclass(frozen=True)
class SchemaTemplateResult:
    output_path: Path
    table_name: str
    sheet_name: str
    column_count: int


@dataclass(frozen=True)
class WorkbookTemplateResult:
    output_path: Path
    sheet_names: tuple[str, ...]
    table_count: int


class LinkValidationWarning(NamedTuple):
    entity_type: str
    column: str
    target_key: str
    unmatched_count: int


class LoadSchemaError(ValueError):
    """Raised when a load schema is invalid."""


class SpreadsheetLoadError(ValueError):
    """Raised when a spreadsheet cannot be loaded with the provided schema."""


def copy_pod_schema_template(
    output_path: Path | str | None = None,
    overwrite: bool = False,
) -> Path:
    """Create the built-in POD workbook template at a user-visible path."""
    return (
        generate_pod_workbook_template(
            output_path=output_path, overwrite=overwrite
        ).output_path
    )


@dataclass(frozen=True)
class ColumnSchema:
    name: str
    type: str
    description: str
    aliases: tuple[str, ...] = field(default_factory=tuple)
    unit: str | None = None
    maps_to: tuple[str, ...] = field(default_factory=tuple)

    @property
    def duckdb_type(self) -> str:
        return _TYPE_MAP[self.type]


@dataclass(frozen=True)
class LoadSchema:
    table_name: str
    description: str
    sheet_name: str
    columns: tuple[ColumnSchema, ...]
    links: tuple[LinkSchema, ...]
    raw: dict[str, Any]


@dataclass(frozen=True)
class LoadResult:
    table_name: str
    row_count: int
    column_count: int
    db_path: Path
    link_warnings: tuple[LinkValidationWarning, ...] = ()


@dataclass(frozen=True)
class LinkSchema:
    entity_type: str
    column: str
    target_key: str


@dataclass(frozen=True)
class WorkbookSheetSpec:
    table_name: str
    description: str
    columns: tuple[ColumnSchema, ...]
    links: tuple[LinkSchema, ...] = ()


def _validate_identifier(value: Any, field_name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise LoadSchemaError(f"{field_name} must be a non-empty string.")
    value = value.strip()
    if not _IDENTIFIER_RE.fullmatch(value):
        raise LoadSchemaError(
            f"{field_name} must use only letters, numbers, and underscores, "
            "and must not start with a number."
        )
    return value


def _validate_text(value: Any, field_name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise LoadSchemaError(f"{field_name} must be a non-empty string.")
    return value.strip()


def _sanitize_identifier(value: str) -> str:
    identifier = re.sub(r"[^A-Za-z0-9_]+", "_", value.strip().lower())
    identifier = re.sub(r"_+", "_", identifier).strip("_")
    if not identifier:
        identifier = "loaded_table"
    if identifier[0].isdigit():
        identifier = f"table_{identifier}"
    return identifier


def _infer_schema_type(series: pd.Series) -> str:
    dtype = series.dropna().dtype
    if pd.api.types.is_bool_dtype(dtype):
        return "boolean"
    if pd.api.types.is_integer_dtype(dtype):
        return "integer"
    if pd.api.types.is_float_dtype(dtype):
        return "double"
    if pd.api.types.is_datetime64_any_dtype(dtype):
        non_null = series.dropna()
        if not non_null.empty:
            normalized = non_null.dt.normalize()
            if (non_null == normalized).all():
                return "date"
        return "datetime"
    return "string"


def generate_schema_template_from_excel(
    excel_path: Path | str,
    output_path: Path | str | None = None,
    overwrite: bool = False,
) -> SchemaTemplateResult:
    """Generate a starter YAML schema from the first Excel sheet."""
    excel = Path(excel_path)
    if not excel.exists():
        raise SpreadsheetLoadError(f"Excel file not found: {excel}")
    if excel.suffix.lower() != ".xlsx":
        raise SpreadsheetLoadError("Only .xlsx Excel files are supported.")

    destination = Path(output_path) if output_path else Path.cwd() / (
        f"{excel.stem}.schema.yaml"
    )
    if destination.exists() and not overwrite:
        raise SpreadsheetLoadError(
            f"Output schema already exists: {destination}. "
            "Use --overwrite to replace it."
        )

    try:
        workbook = pd.ExcelFile(excel, engine="openpyxl")
        sheet_name = workbook.sheet_names[0]
        df = pd.read_excel(workbook, sheet_name=sheet_name)
    except Exception as e:
        raise SpreadsheetLoadError(f"Failed to read Excel file: {e}") from e

    columns: list[dict[str, Any]] = []
    links: list[dict[str, str]] = []
    seen: set[str] = set()
    for raw_column in df.columns:
        name = str(raw_column).strip()
        name = _validate_identifier(name, f"Excel column '{raw_column}'")
        if name in seen:
            raise SpreadsheetLoadError(f"Duplicate Excel column name: {name}")
        seen.add(name)
        columns.append(
            {
                "name": name,
                "type": _infer_schema_type(df[name]),
                "description": f"TODO: describe {name}",
                "aliases": [],
            }
        )
        for entity_type, target_key in _LINK_TARGET_KEYS.items():
            if name == target_key:
                links.append(
                    {
                        "entity_type": entity_type,
                        "column": name,
                        "target_key": target_key,
                    }
                )

    if not columns:
        raise SpreadsheetLoadError("Excel sheet must contain at least one column.")

    schema = {
        "table_name": _sanitize_identifier(excel.stem),
        "description": "TODO: describe this table",
        "sheet_name": sheet_name,
        "columns": columns,
    }
    if links:
        schema["links"] = links
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        yaml.safe_dump(schema, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    return SchemaTemplateResult(
        output_path=destination,
        table_name=schema["table_name"],
        sheet_name=sheet_name,
        column_count=len(columns),
    )


def load_schema_from_yaml(schema_path: Path | str) -> LoadSchema:
    """Load and validate an `esdc load` YAML schema."""
    path = Path(schema_path)
    if not path.exists():
        raise LoadSchemaError(f"Schema file not found: {path}")
    try:
        with open(path, encoding="utf-8") as f:
            raw = yaml.safe_load(f)
    except yaml.YAMLError as e:
        raise LoadSchemaError(f"Schema YAML is invalid: {e}") from e
    if not isinstance(raw, dict):
        raise LoadSchemaError("Schema YAML must be a mapping.")

    table_name = _validate_identifier(raw.get("table_name"), "table_name")
    description = _validate_text(raw.get("description"), "description")
    sheet_name = _validate_text(raw.get("sheet_name"), "sheet_name")
    raw_columns = raw.get("columns")
    if not isinstance(raw_columns, list) or not raw_columns:
        raise LoadSchemaError("columns must be a non-empty list.")

    columns: list[ColumnSchema] = []
    seen: set[str] = set()
    for idx, item in enumerate(raw_columns, start=1):
        if not isinstance(item, dict):
            raise LoadSchemaError(f"columns[{idx}] must be a mapping.")
        name = _validate_identifier(item.get("name"), f"columns[{idx}].name")
        if name in seen:
            raise LoadSchemaError(f"Duplicate column name: {name}")
        seen.add(name)
        type_name = _validate_text(item.get("type"), f"columns[{idx}].type").lower()
        if type_name not in _TYPE_MAP:
            valid = ", ".join(sorted(_TYPE_MAP))
            raise LoadSchemaError(
                f"Unsupported column type '{type_name}'. Use: {valid}"
            )
        col_description = _validate_text(
            item.get("description"), f"columns[{idx}].description"
        )
        raw_aliases = item.get("aliases", []) or []
        if not isinstance(raw_aliases, list) or not all(
            isinstance(alias, str) for alias in raw_aliases
        ):
            raise LoadSchemaError(f"columns[{idx}].aliases must be a list of strings.")
        raw_unit = item.get("unit")
        unit = str(raw_unit).strip() if raw_unit is not None else None
        raw_maps_to = item.get("maps_to", []) or []
        if not isinstance(raw_maps_to, list) or not all(
            isinstance(target, str) for target in raw_maps_to
        ):
            raise LoadSchemaError(f"columns[{idx}].maps_to must be a list of strings.")
        columns.append(
            ColumnSchema(
                name=name,
                type=type_name,
                description=col_description,
                aliases=tuple(alias.strip() for alias in raw_aliases if alias.strip()),
                unit=unit or None,
                maps_to=tuple(
                    target.strip() for target in raw_maps_to if target.strip()
                ),
            )
        )

    column_names = {column.name for column in columns}
    links: list[LinkSchema] = []
    raw_links = raw.get("links", []) or []
    if not isinstance(raw_links, list):
        raise LoadSchemaError("links must be a list.")
    for idx, item in enumerate(raw_links, start=1):
        if not isinstance(item, dict):
            raise LoadSchemaError(f"links[{idx}] must be a mapping.")
        entity_type = _validate_text(
            item.get("entity_type"), f"links[{idx}].entity_type"
        )
        if entity_type not in _LINK_TARGET_KEYS:
            valid = ", ".join(_LINK_TARGET_KEYS)
            raise LoadSchemaError(
                f"Unsupported links[{idx}].entity_type '{entity_type}'. Use: {valid}"
            )
        column = _validate_identifier(item.get("column"), f"links[{idx}].column")
        if column not in column_names:
            raise LoadSchemaError(f"links[{idx}].column '{column}' is not in columns.")
        target_key = item.get("target_key") or _LINK_TARGET_KEYS[entity_type]
        target_key = _validate_identifier(target_key, f"links[{idx}].target_key")
        links.append(
            LinkSchema(entity_type=entity_type, column=column, target_key=target_key)
        )

    normalized_columns: list[dict[str, Any]] = []
    for column in columns:
        item: dict[str, Any] = {
            "name": column.name,
            "type": column.type,
            "description": column.description,
            "aliases": list(column.aliases),
        }
        if column.unit:
            item["unit"] = column.unit
        if column.maps_to:
            item["maps_to"] = list(column.maps_to)
        normalized_columns.append(item)
    normalized_raw: dict[str, Any] = {
        "table_name": table_name,
        "description": description,
        "sheet_name": sheet_name,
        "columns": normalized_columns,
    }
    if links:
        normalized_raw["links"] = [
            {
                "entity_type": link.entity_type,
                "column": link.column,
                "target_key": link.target_key,
            }
            for link in links
        ]

    return LoadSchema(
        table_name=table_name,
        description=description,
        sheet_name=sheet_name,
        columns=tuple(columns),
        links=tuple(links),
        raw=normalized_raw,
    )


def _load_pod_domain_schema() -> LoadSchema:
    if not POD_SCHEMA_PATH.exists():
        raise LoadSchemaError(f"POD schema template not found: {POD_SCHEMA_PATH}")
    return load_schema_from_yaml(POD_SCHEMA_PATH)


def _pod_metric_columns(schema: LoadSchema) -> tuple[ColumnSchema, ...]:
    excluded = {
        "report_date",
        "effective_date",
        "case_type",
        "pod_id",
        "pod_letter_num",
        "pod_name",
        "pod_scope",
        "supercedes_by",
    }
    return tuple(column for column in schema.columns if column.name not in excluded)


def _build_pod_workbook_specs() -> tuple[WorkbookSheetSpec, ...]:
    schema = _load_pod_domain_schema()
    column_by_name = {column.name: column for column in schema.columns}
    metric_columns = _pod_metric_columns(schema)

    plan_columns: tuple[ColumnSchema, ...] = tuple(
        column_by_name[name]
        for name in (
            "pod_id",
            "pod_letter_num",
            "pod_name",
            "pod_scope",
            "supercedes_by",
            "report_date",
            "effective_date",
        )
        if name in column_by_name
    ) + metric_columns

    monitoring_columns: tuple[ColumnSchema, ...] = (
        column_by_name["pod_id"],
        column_by_name["case_type"],
        column_by_name["report_date"],
        column_by_name["effective_date"],
    ) + metric_columns

    project_columns = (
        ColumnSchema(
            name="pod_id",
            type="string",
            description="Unique POD identifier.",
            aliases=("ID POD", "POD ID"),
            maps_to=("ProducingLicense",),
        ),
        ColumnSchema(
            name="project_id",
            type="string",
            description="Linked project identifier.",
            aliases=("project code", "kode proyek"),
        ),
    )

    return (
        WorkbookSheetSpec(
            table_name=POD_PLAN_SHEET_NAME,
            description="Baseline POD plan data.",
            columns=plan_columns,
        ),
        WorkbookSheetSpec(
            table_name=POD_PROJECT_SHEET_NAME,
            description="Many-to-many POD to project mapping.",
            columns=project_columns,
            links=(
                LinkSchema(
                    entity_type="Project",
                    column="project_id",
                    target_key="project_id",
                ),
            ),
        ),
        WorkbookSheetSpec(
            table_name=POD_MONITORING_SHEET_NAME,
            description="POD outlook and actual monitoring data.",
            columns=monitoring_columns,
        ),
    )


def _sheet_spec_to_load_schema(spec: WorkbookSheetSpec) -> LoadSchema:
    raw_columns: list[dict[str, Any]] = []
    for column in spec.columns:
        item: dict[str, Any] = {
            "name": column.name,
            "type": column.type,
            "description": column.description,
            "aliases": list(column.aliases),
        }
        if column.unit:
            item["unit"] = column.unit
        if column.maps_to:
            item["maps_to"] = list(column.maps_to)
        raw_columns.append(item)

    return LoadSchema(
        table_name=spec.table_name,
        description=spec.description,
        sheet_name=spec.table_name,
        columns=spec.columns,
        links=spec.links,
        raw={
            "table_name": spec.table_name,
            "description": spec.description,
            "sheet_name": spec.table_name,
            "columns": raw_columns,
            "links": [
                {
                    "entity_type": link.entity_type,
                    "column": link.column,
                    "target_key": link.target_key,
                }
                for link in spec.links
            ],
        },
    )


def _write_sheet_header(
    ws: Any,
    spec: WorkbookSheetSpec,
    *,
    add_case_type_validation: bool = False,
) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(vertical="center", horizontal="center")
    description_row = 1
    header_row = 2
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[description_row].height = 24
    ws.row_dimensions[header_row].height = 24

    ws.merge_cells(
        start_row=description_row,
        start_column=1,
        end_row=description_row,
        end_column=max(len(spec.columns), 1),
    )
    ws.cell(description_row, 1).value = spec.description
    ws.cell(description_row, 1).font = Font(bold=True, color="1F1F1F")
    ws.cell(description_row, 1).alignment = Alignment(horizontal="left")

    for idx, column in enumerate(spec.columns, start=1):
        cell = ws.cell(header_row, idx)
        cell.value = column.name
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        comment_parts = [f"Type: {column.type}", f"Description: {column.description}"]
        if column.unit:
            comment_parts.append(f"Unit: {column.unit}")
        if column.aliases:
            comment_parts.append(f"Aliases: {', '.join(column.aliases)}")
        if column.maps_to:
            comment_parts.append(f"Maps to: {', '.join(column.maps_to)}")
        cell.comment = Comment("\n".join(comment_parts), "Codex")
        ws.column_dimensions[get_column_letter(idx)].width = max(
            14,
            min(34, len(column.name) + 4),
        )

    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(spec.columns))}{header_row}"
    )

    if add_case_type_validation:
        from openpyxl.worksheet.datavalidation import DataValidation

        dv = DataValidation(
            type="list",
            formula1='"outlook,actual"',
            allow_blank=False,
        )
        ws.add_data_validation(dv)
        dv.add("B3:B1048576")


def generate_pod_workbook_template(
    output_path: Path | str | None = None,
    overwrite: bool = False,
) -> WorkbookTemplateResult:
    """Generate the built-in POD workbook template."""
    destination = (
        Path(output_path) if output_path else Path.cwd() / POD_TEMPLATE_FILENAME
    )
    if destination.exists() and not overwrite:
        raise SpreadsheetLoadError(
            f"Output schema already exists: {destination}. "
            "Use --overwrite to replace it."
        )

    workbook = Workbook()
    workbook.remove(workbook.active)
    specs = _build_pod_workbook_specs()
    for spec in specs:
        ws = workbook.create_sheet(title=spec.table_name)
        _write_sheet_header(
            ws,
            spec,
            add_case_type_validation=spec.table_name == POD_MONITORING_SHEET_NAME,
        )

    from openpyxl.worksheet.datavalidation import DataValidation

    dn = DefinedName("pod_plan_ids", attr_text=f"{POD_PLAN_SHEET_NAME}!$A$3:$A$1048576")
    workbook.defined_names.add(dn)

    for sheet_name in (POD_PROJECT_SHEET_NAME, POD_MONITORING_SHEET_NAME):
        dv = DataValidation(
            type="list",
            formula1='INDIRECT("pod_plan_ids")',
            allow_blank=True,
            error="Pilih POD ID yang valid dari sheet pod_plan",
            errorTitle="POD ID tidak valid",
        )
        workbook[sheet_name].add_data_validation(dv)
        dv.add("A3:A1048576")

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return WorkbookTemplateResult(
        output_path=destination,
        sheet_names=tuple(spec.table_name for spec in specs),
        table_count=len(specs),
    )


def _create_metadata_table(conn: duckdb.DuckDBPyConnection) -> None:
    conn.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {_METADATA_TABLE} (
            table_name TEXT PRIMARY KEY,
            description TEXT,
            sheet_name TEXT,
            schema_yaml TEXT,
            loaded_at TEXT
        )
        """
    )


def _coerce_dataframe(df: pd.DataFrame, schema: LoadSchema) -> pd.DataFrame:
    expected = [column.name for column in schema.columns]
    actual = [str(column) for column in df.columns]
    missing = [column for column in expected if column not in actual]
    extra = [column for column in actual if column not in expected]
    if missing or extra:
        parts: list[str] = []
        if missing:
            parts.append(f"missing columns: {', '.join(missing)}")
        if extra:
            parts.append(f"unexpected columns: {', '.join(extra)}")
        raise SpreadsheetLoadError(
            "Excel columns do not match schema: " + "; ".join(parts)
        )

    coerced = df.loc[:, expected].copy()
    for column in schema.columns:
        try:
            if column.type in {"string", "text"}:
                coerced[column.name] = coerced[column.name].astype("string")
            elif column.type in {"integer", "int", "bigint"}:
                coerced[column.name] = pd.to_numeric(
                    coerced[column.name], errors="raise"
                ).astype("Int64")
            elif column.type in {"double", "float", "number"}:
                coerced[column.name] = pd.to_numeric(
                    coerced[column.name], errors="raise"
                )
            elif column.type in {"boolean", "bool"}:
                coerced[column.name] = coerced[column.name].astype("boolean")
            elif column.type == "date":
                coerced[column.name] = pd.to_datetime(
                    coerced[column.name], errors="raise"
                ).dt.date
            elif column.type == "datetime":
                coerced[column.name] = pd.to_datetime(
                    coerced[column.name], errors="raise"
                )
        except (TypeError, ValueError) as e:
            raise SpreadsheetLoadError(
                f"Column '{column.name}' cannot be converted to {column.type}: {e}"
            ) from e
    return coerced


def _table_has_column(
    conn: duckdb.DuckDBPyConnection,
    table_name: str,
    column_name: str,
) -> bool:
    row = conn.execute(
        """
        SELECT COUNT(*)
        FROM information_schema.columns
        WHERE table_name = ? AND column_name = ?
        """,
        [table_name, column_name],
    ).fetchone()
    return bool(row and row[0] > 0)


def _validate_link_values(
    conn: duckdb.DuckDBPyConnection,
    df: pd.DataFrame,
    schema: LoadSchema,
) -> tuple[LinkValidationWarning, ...]:
    warnings: list[LinkValidationWarning] = []
    for link in schema.links:
        source_values = {
            str(value)
            for value in df[link.column].dropna().unique().tolist()
            if str(value).strip()
        }
        if not source_values:
            continue

        reference_values: set[str] = set()
        for table_name in _LINK_REFERENCE_TABLES:
            if not _table_has_column(conn, table_name, link.target_key):
                continue
            rows = conn.execute(
                f"SELECT DISTINCT {link.target_key} FROM {table_name}"
            ).fetchall()
            reference_values.update(str(row[0]) for row in rows if row[0] is not None)

        if not reference_values:
            continue
        unmatched = source_values - reference_values
        if unmatched:
            warnings.append(
                LinkValidationWarning(
                    entity_type=link.entity_type,
                    column=link.column,
                    target_key=link.target_key,
                    unmatched_count=len(unmatched),
                )
            )
    return tuple(warnings)


def _write_dataframe_to_table(
    conn: duckdb.DuckDBPyConnection,
    schema: LoadSchema,
    df: pd.DataFrame,
) -> tuple[int, tuple[LinkValidationWarning, ...]]:
    columns_sql = ", ".join(
        f"{column.name} {column.duckdb_type}" for column in schema.columns
    )
    conn.execute(f"DROP TABLE IF EXISTS {schema.table_name}")
    conn.execute(f"CREATE TABLE {schema.table_name} ({columns_sql})")
    conn.register("_esdc_load_df", df)
    try:
        conn.execute(f"INSERT INTO {schema.table_name} SELECT * FROM _esdc_load_df")
    finally:
        with contextlib.suppress(Exception):
            conn.unregister("_esdc_load_df")
    link_warnings = _validate_link_values(conn, df, schema)
    return len(df), link_warnings


def _load_excel_sheet(
    excel: Path,
    sheet_name: str,
    schema: LoadSchema,
    *,
    header_row: int = 0,
) -> pd.DataFrame:
    try:
        df = pd.read_excel(
            excel,
            sheet_name=sheet_name,
            engine="openpyxl",
            header=header_row,
        )
    except ValueError as e:
        raise SpreadsheetLoadError(
            f"Excel sheet '{sheet_name}' not found in {excel}."
        ) from e
    except Exception as e:
        raise SpreadsheetLoadError(f"Failed to read Excel file: {e}") from e
    return _coerce_dataframe(df, schema)


def _validate_pod_workbook_crosslinks(
    plan_df: pd.DataFrame,
    project_df: pd.DataFrame,
    monitoring_df: pd.DataFrame,
) -> None:
    if plan_df["pod_id"].duplicated().any():
        raise SpreadsheetLoadError("pod_plan contains duplicate pod_id values.")

    plan_ids = {
        str(value).strip()
        for value in plan_df["pod_id"].dropna().tolist()
        if str(value).strip()
    }

    project_missing = {
        str(value).strip()
        for value in project_df["pod_id"].dropna().tolist()
        if str(value).strip()
    } - plan_ids
    if project_missing:
        raise SpreadsheetLoadError(
            "pod_project references pod_id values not found in pod_plan: "
            f"{', '.join(sorted(project_missing))}"
        )

    monitoring_missing = {
        str(value).strip()
        for value in monitoring_df["pod_id"].dropna().tolist()
        if str(value).strip()
    } - plan_ids
    if monitoring_missing:
        raise SpreadsheetLoadError(
            "pod_monitoring references pod_id values not found in pod_plan: "
            f"{', '.join(sorted(monitoring_missing))}"
        )

    invalid_case_types = {
        str(value).strip()
        for value in monitoring_df["case_type"].dropna().tolist()
        if str(value).strip()
    } - {"outlook", "actual"}
    if invalid_case_types:
        raise SpreadsheetLoadError(
            "pod_monitoring.case_type must be one of outlook, actual. "
            f"Invalid values: {', '.join(sorted(invalid_case_types))}"
        )

    if project_df.duplicated(subset=["pod_id", "project_id"]).any():
        raise SpreadsheetLoadError(
            "pod_project contains duplicate pod_id/project_id pairs."
        )

    if monitoring_df.duplicated(subset=["pod_id", "case_type", "effective_date"]).any():
        raise SpreadsheetLoadError(
            "pod_monitoring contains duplicate pod_id/case_type/effective_date rows."
        )


def _create_pod_views(conn: duckdb.DuckDBPyConnection) -> None:
    schema = _load_pod_domain_schema()
    metric_names = tuple(
        column.name for column in _pod_metric_columns(schema)
    )
    plan_meta = ["pod_letter_num", "pod_name", "pod_scope", "supercedes_by"]

    metrics_sql = ", ".join(metric_names)
    plan_meta_sql = ", ".join(plan_meta)
    null_meta_sql = ", ".join(f"NULL::TEXT AS {name}" for name in plan_meta)

    conn.execute(f"""
        CREATE OR REPLACE VIEW pod_economics AS
        SELECT
            pod_id,
            'plan'::TEXT AS case_type,
            report_date,
            effective_date,
            {plan_meta_sql},
            {metrics_sql}
        FROM pod_plan
        UNION ALL
        SELECT
            pod_id,
            case_type,
            report_date,
            effective_date,
            {null_meta_sql},
            {metrics_sql}
        FROM pod_monitoring
    """)

    if _table_has_column(conn, "project_resources", "project_id"):
        conn.execute("""
            CREATE OR REPLACE VIEW pod_project_economics AS
            WITH ranked_pr AS (
                SELECT *,
                    ROW_NUMBER() OVER (
                        PARTITION BY project_id
                        ORDER BY report_date DESC NULLS LAST
                    ) AS rn
                FROM project_resources
            )
            SELECT
                pe.*,
                pp.project_id,
                rpr.project_name,
                rpr.project_stage,
                rpr.project_class,
                rpr.field_name,
                rpr.wk_name
            FROM pod_economics pe
            JOIN pod_project pp ON pe.pod_id = pp.pod_id
            LEFT JOIN ranked_pr rpr ON pp.project_id = rpr.project_id AND rpr.rn = 1
        """)


def load_pod_workbook_to_duckdb(excel_path: Path | str) -> tuple[LoadResult, ...]:
    """Load the built-in POD workbook into DuckDB tables."""
    excel = Path(excel_path)
    if not excel.exists():
        raise SpreadsheetLoadError(f"Excel file not found: {excel}")
    if excel.suffix.lower() != ".xlsx":
        raise SpreadsheetLoadError("Only .xlsx Excel files are supported.")

    specs = _build_pod_workbook_specs()
    schemas = {spec.table_name: _sheet_spec_to_load_schema(spec) for spec in specs}

    plan_schema = schemas[POD_PLAN_SHEET_NAME]
    project_schema = schemas[POD_PROJECT_SHEET_NAME]
    monitoring_schema = schemas[POD_MONITORING_SHEET_NAME]

    plan_df = _load_excel_sheet(
        excel, POD_PLAN_SHEET_NAME, plan_schema, header_row=1
    )
    project_df = _load_excel_sheet(
        excel, POD_PROJECT_SHEET_NAME, project_schema, header_row=1
    )
    monitoring_df = _load_excel_sheet(
        excel, POD_MONITORING_SHEET_NAME, monitoring_schema, header_row=1
    )
    _validate_pod_workbook_crosslinks(plan_df, project_df, monitoring_df)

    db_path = Config.get_db_file()
    _ensure_duckdb_database(db_path)
    Config.get_db_dir().mkdir(parents=True, exist_ok=True)
    conn = get_duckdb_connection(db_path)
    results: list[LoadResult] = []
    try:
        conn.execute("BEGIN")
        for schema, df in (
            (plan_schema, plan_df),
            (project_schema, project_df),
            (monitoring_schema, monitoring_df),
        ):
            row_count, link_warnings = _write_dataframe_to_table(conn, schema, df)
            _create_metadata_table(conn)
            schema_yaml = yaml.safe_dump(
                schema.raw, sort_keys=False, allow_unicode=True
            )
            conn.execute(
                f"DELETE FROM {_METADATA_TABLE} WHERE table_name = ?",
                [schema.table_name],
            )
            conn.execute(
                f"""
                INSERT INTO {_METADATA_TABLE}
                    (table_name, description, sheet_name, schema_yaml, loaded_at)
                VALUES (?, ?, ?, ?, ?)
                """,
                [
                    schema.table_name,
                    schema.description,
                    schema.sheet_name,
                    schema_yaml,
                    datetime.now(timezone.utc).isoformat(),
                ],
            )
            results.append(
                LoadResult(
                    table_name=schema.table_name,
                    row_count=row_count,
                    column_count=len(schema.columns),
                    db_path=db_path,
                    link_warnings=link_warnings,
                )
            )
        _create_pod_views(conn)
        conn.execute("COMMIT")
        conn.execute("CHECKPOINT")
    except Exception:
        with contextlib.suppress(Exception):
            conn.execute("ROLLBACK")
        raise
    finally:
        conn.close()

    from esdc.chat.tools import invalidate_tool_cache, reset_sql_cache

    reset_sql_cache()
    invalidate_tool_cache()
    return tuple(results)


def _cypher_escape(value: Any) -> str:
    return str(value).replace("\\", "\\\\").replace("'", "\\'")


def build_loaded_schema_graph(rows: list[tuple[str, str, str, str]]) -> Any:
    """Build a LadybugDB metadata graph for loaded spreadsheet schemas."""
    import real_ladybug as lb

    db = lb.Database(":memory:")
    conn = lb.Connection(db)
    statements = [
        (
            "CREATE NODE TABLE LoadedTable "
            "(table_name STRING PRIMARY KEY, description STRING, sheet_name STRING)"
        ),
        (
            "CREATE NODE TABLE LoadedColumn "
            "(column_id STRING PRIMARY KEY, table_name STRING, name STRING, "
            "description STRING, data_type STRING, aliases STRING, unit STRING)"
        ),
        (
            "CREATE NODE TABLE ReferenceEntity "
            "(entity_type STRING PRIMARY KEY, target_key STRING)"
        ),
        "CREATE NODE TABLE KSMIConcept (code STRING PRIMARY KEY)",
        (
            "CREATE REL TABLE LOADED_TABLE_HAS_COLUMN "
            "(FROM LoadedTable TO LoadedColumn)"
        ),
        (
            "CREATE REL TABLE LOADED_TABLE_LINKS_TO_ENTITY "
            "(FROM LoadedTable TO ReferenceEntity, "
            "link_column STRING, target_key STRING)"
        ),
        (
            "CREATE REL TABLE LOADED_COLUMN_MAPS_TO_KSMI "
            "(FROM LoadedColumn TO KSMIConcept)"
        ),
    ]
    for stmt in statements:
        conn.execute(stmt)

    reference_entities: set[str] = set()
    ksmi_concepts: set[str] = set()
    column_ids: list[tuple[str, str]] = []
    table_links: list[tuple[str, str, str, str]] = []
    column_maps: list[tuple[str, str]] = []

    for table_name, description, sheet_name, schema_yaml in rows:
        table = _cypher_escape(table_name)
        conn.execute(
            f"CREATE (t:LoadedTable "
            f"{{table_name: '{table}', "
            f"description: '{_cypher_escape(description)}', "
            f"sheet_name: '{_cypher_escape(sheet_name)}'}})"
        )
        data = yaml.safe_load(schema_yaml) or {}
        for column in data.get("columns", []):
            if not isinstance(column, dict):
                continue
            column_name = str(column.get("name", ""))
            column_id = f"{table_name}.{column_name}"
            column_ids.append((table_name, column_id))
            aliases = "||".join(column.get("aliases", []) or [])
            conn.execute(
                f"CREATE (c:LoadedColumn "
                f"{{column_id: '{_cypher_escape(column_id)}', "
                f"table_name: '{table}', "
                f"name: '{_cypher_escape(column_name)}', "
                f"description: '{_cypher_escape(column.get('description', ''))}', "
                f"data_type: '{_cypher_escape(column.get('type', ''))}', "
                f"aliases: '{_cypher_escape(aliases)}', "
                f"unit: '{_cypher_escape(column.get('unit', '') or '')}'}})"
            )
            for concept in column.get("maps_to", []) or []:
                concept_code = str(concept)
                ksmi_concepts.add(concept_code)
                column_maps.append((column_id, concept_code))

        for link in data.get("links", []) or []:
            if not isinstance(link, dict):
                continue
            entity_type = str(link.get("entity_type", ""))
            target_key = str(link.get("target_key", ""))
            reference_entities.add(entity_type)
            table_links.append(
                (table_name, entity_type, str(link.get("column", "")), target_key)
            )

    for entity_type in reference_entities:
        target_key = _LINK_TARGET_KEYS.get(entity_type, "")
        conn.execute(
            f"CREATE (r:ReferenceEntity "
            f"{{entity_type: '{_cypher_escape(entity_type)}', "
            f"target_key: '{_cypher_escape(target_key)}'}})"
        )
    for concept_code in ksmi_concepts:
        conn.execute(
            f"CREATE (k:KSMIConcept "
            f"{{code: '{_cypher_escape(concept_code)}'}})"
        )

    for table_name, column_id in column_ids:
        conn.execute(
            f"MATCH (t:LoadedTable "
            f"{{table_name: '{_cypher_escape(table_name)}'}}), "
            f"(c:LoadedColumn "
            f"{{column_id: '{_cypher_escape(column_id)}'}}) "
            f"CREATE (t)-[:LOADED_TABLE_HAS_COLUMN]->(c)"
        )
    for table_name, entity_type, column_name, target_key in table_links:
        conn.execute(
            f"MATCH (t:LoadedTable "
            f"{{table_name: '{_cypher_escape(table_name)}'}}), "
            f"(r:ReferenceEntity "
            f"{{entity_type: '{_cypher_escape(entity_type)}'}}) "
            f"CREATE (t)-[:LOADED_TABLE_LINKS_TO_ENTITY "
            f"{{link_column: '{_cypher_escape(column_name)}', "
            f"target_key: '{_cypher_escape(target_key)}'}}]->(r)"
        )
    for column_id, concept_code in column_maps:
        conn.execute(
            f"MATCH (c:LoadedColumn "
            f"{{column_id: '{_cypher_escape(column_id)}'}}), "
            f"(k:KSMIConcept {{code: '{_cypher_escape(concept_code)}'}}) "
            f"CREATE (c)-[:LOADED_COLUMN_MAPS_TO_KSMI]->(k)"
        )
    return conn


def load_excel_to_duckdb(
    excel_path: Path | str,
    schema_path: Path | str,
) -> LoadResult:
    """Load an Excel sheet into a replace-mode DuckDB table."""
    schema = load_schema_from_yaml(schema_path)
    excel = Path(excel_path)
    if not excel.exists():
        raise SpreadsheetLoadError(f"Excel file not found: {excel}")
    if excel.suffix.lower() != ".xlsx":
        raise SpreadsheetLoadError("Only .xlsx Excel files are supported.")

    try:
        df = pd.read_excel(excel, sheet_name=schema.sheet_name, engine="openpyxl")
    except ValueError as e:
        raise SpreadsheetLoadError(
            f"Excel sheet '{schema.sheet_name}' not found in {excel}."
        ) from e
    except Exception as e:
        raise SpreadsheetLoadError(f"Failed to read Excel file: {e}") from e

    df = _coerce_dataframe(df, schema)
    db_path = Config.get_db_file()
    _ensure_duckdb_database(db_path)
    Config.get_db_dir().mkdir(parents=True, exist_ok=True)

    conn = get_duckdb_connection(db_path)
    link_warnings: tuple[LinkValidationWarning, ...] = ()
    row_count = len(df)
    try:
        conn.execute("BEGIN")
        row_count, link_warnings = _write_dataframe_to_table(conn, schema, df)

        _create_metadata_table(conn)
        schema_yaml = yaml.safe_dump(schema.raw, sort_keys=False, allow_unicode=True)
        conn.execute(
            f"DELETE FROM {_METADATA_TABLE} WHERE table_name = ?",
            [schema.table_name],
        )
        conn.execute(
            f"""
            INSERT INTO {_METADATA_TABLE}
                (table_name, description, sheet_name, schema_yaml, loaded_at)
            VALUES (?, ?, ?, ?, ?)
            """,
            [
                schema.table_name,
                schema.description,
                schema.sheet_name,
                schema_yaml,
                datetime.now(timezone.utc).isoformat(),
            ],
        )
        conn.execute("COMMIT")
        conn.execute("CHECKPOINT")
    except Exception:
        with contextlib.suppress(Exception):
            conn.execute("ROLLBACK")
        raise
    finally:
        conn.close()

    from esdc.chat.tools import invalidate_tool_cache, reset_sql_cache

    reset_sql_cache()
    invalidate_tool_cache()
    return LoadResult(
        table_name=schema.table_name,
        row_count=row_count,
        column_count=len(schema.columns),
        db_path=db_path,
        link_warnings=link_warnings,
    )


def lookup_loaded_schema(entity: str | None) -> str | None:
    """Return data dictionary text for a loaded table or column entity."""
    if not entity or not Config.get_db_file().exists():
        return None
    try:
        conn = get_duckdb_connection(Config.get_db_file(), read_only=True)
    except Exception:
        return None

    needle = entity.casefold()
    try:
        exists = conn.execute(
            "SELECT COUNT(*) FROM information_schema.tables WHERE table_name = ?",
            [_METADATA_TABLE],
        ).fetchone()
        if not exists or exists[0] == 0:
            return None
        rows = conn.execute(
            "SELECT table_name, description, sheet_name, schema_yaml "
            f"FROM {_METADATA_TABLE}"
        ).fetchall()
    finally:
        conn.close()

    with contextlib.suppress(Exception):
        build_loaded_schema_graph(rows)

    for table_name, description, sheet_name, schema_yaml in rows:
        data = yaml.safe_load(schema_yaml) or {}
        if str(table_name).casefold() == needle:
            columns = data.get("columns", [])
            links = data.get("links", []) or []
            lines = [
                f"Loaded table: {table_name}",
                f"Description: {description}",
                f"Excel sheet: {sheet_name}",
                "Columns:",
            ]
            for column in columns:
                if isinstance(column, dict):
                    unit = f" ({column['unit']})" if column.get("unit") else ""
                    lines.append(
                        f"- {column.get('name')}: {column.get('description')} "
                        f"[{column.get('type')}{unit}]"
                    )
            if links:
                lines.append("Entity links:")
                for link in links:
                    if isinstance(link, dict):
                        lines.append(
                            f"- {link.get('column')} links to "
                            f"{link.get('entity_type')}.{link.get('target_key')}"
                        )
            return "\n".join(lines)

        for column in data.get("columns", []):
            if not isinstance(column, dict):
                continue
            aliases = column.get("aliases", []) or []
            names = [column.get("name"), *aliases]
            match = any(
                isinstance(name, str) and name.casefold() == needle for name in names
            )
            if match:
                unit = f"\nUnit: {column['unit']}" if column.get("unit") else ""
                alias_text = (
                    f"\nAliases: {', '.join(aliases)}" if aliases else ""
                )
                maps_to = column.get("maps_to", []) or []
                maps_to_text = (
                    f"\nMaps to KSMI: {', '.join(maps_to)}" if maps_to else ""
                )
                return (
                    f"Loaded column: {table_name}.{column.get('name')}\n"
                    f"Description: {column.get('description')}\n"
                    f"Type: {column.get('type')}"
                    f"{unit}"
                    f"{alias_text}\n"
                    f"{maps_to_text}"
                    "\n"
                    "Source: loaded spreadsheet data dictionary."
                )
    return None


def print_load_result(result: LoadResult) -> None:
    """Print a concise load summary for the CLI."""
    typer.echo(
        f"Loaded {result.row_count:,} rows x "
        f"{result.column_count:,} columns into '{result.table_name}'"
    )
    for warning in result.link_warnings:
        typer.echo(
            "Warning: "
            f"{warning.unmatched_count:,} value(s) in '{warning.column}' "
            f"did not match {warning.entity_type}.{warning.target_key}"
        )
    typer.echo(f"Database: {result.db_path}")
